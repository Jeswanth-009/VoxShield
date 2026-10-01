from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\anc")
OUT = ROOT / "docs" / "VoxShield_Master_Technical_and_Presentation_Source.docx"

NAVY = "102A56"
BLUE = "1665D8"
CYAN = "16A6C9"
ORANGE = "F06A21"
GREEN = "17845B"
RED = "C8362E"
LIGHT = "EAF2FA"
PALE = "F6F8FB"
GREY = "5A6472"


def load_json(path: str):
    p = ROOT / path
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(section):
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None, header_fill=NAVY):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = True
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", 8.5)
        shade(t.rows[0].cells[i], header_fill)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, None, 8.2)
            if ridx % 2:
                shade(cells[i], PALE)
    if widths:
        for row in t.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def callout(doc, title, body, fill=LIGHT, accent=BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    r2 = p.add_run(body)
    r2.font.color.rgb = RGBColor.from_string("202833")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def metric(d, *keys, default="—"):
    for k in keys:
        if not isinstance(d, dict) or k not in d:
            return default
        d = d[k]
    return d


def fmt(value, n=3):
    return f"{value:.{n}f}" if isinstance(value, (float, int)) else str(value)


doc = Document()
for s in doc.sections:
    margins(s)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
for name, size, color in [("Title", 32, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 11, GREEN)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)
    styles[name].font.bold = True

section = doc.sections[0]
header = section.header.paragraphs[0]
header.text = "VOXSHIELD  |  MASTER TECHNICAL RECORD & PRESENTATION SOURCE"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(7.5)
header.runs[0].font.color.rgb = RGBColor.from_string(GREY)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("VoxShield • Evidence-based working document • 26 September 2026")
footer.runs[0].font.size = Pt(7.5)
footer.runs[0].font.color.rgb = RGBColor.from_string(GREY)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
r = p.add_run("VOXSHIELD")
r.bold = True; r.font.size = Pt(42); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Hybrid AI/ML Speech Enhancement for Defence Communication")
r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Master Technical Document and Presentation-Content Source")
r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(GREY)
doc.add_paragraph()
table(doc, ["Document status", "Current selected system", "Evidence cut-off"], [["Authoritative project record", "v7 dual-channel fusion UL-UNAS", "26 September 2026"]])
callout(doc, "Purpose", "This is the single main source for the report, judging discussion and PPT. It records what was built, what changed, why changes were necessary, measured results, unresolved limits, feasibility, viability, demo paths and the exact claims that are safe to make.", "EAF2FA", BLUE)
callout(doc, "Evidence discipline", "Measured results are labelled as measured. Targets are labelled as targets. A laptop benchmark is not presented as Raspberry Pi 5 performance. A simulated or laptop-array reference is not presented as a synchronized tactical reference microphone. Speech enhancement is not presented as proven acoustic anti-noise cancellation.", "FFF1E8", ORANGE)
doc.add_page_break()

doc.add_heading("Executive conclusion", level=1)
doc.add_paragraph("VoxShield is now a modular, causal, real-time communication speech-enhancement platform built from an empty ANC project directory. The strongest current artifact is a 172,748-parameter dual-channel fusion UL-UNAS model that uses primary and reference microphone features, estimates a complex ratio mask, exports to stateful ONNX and runs faster than real time on the development laptop. It reaches the SIH absolute SNR, STOI and PESQ targets on the sampled broad validation gate and on the sampled aggregate military gate. It does not yet meet those targets on the hardest military subset, especially helicopter, fighter/aircraft, vehicle and input SNR below −10 dB. Raspberry Pi 5 performance and a correctly synchronized physical two-microphone prototype remain to be measured.")
table(doc, ["Question", "Factual answer"], [
    ["Was the whole project created from scratch?", "Yes. Repository structure, ingestion, mixing, training, evaluation, streaming, export, monitoring, dashboard, tests and deployment scaffolding were created inside C:\\anc."],
    ["Is the current model GTCRN + UL-UNAS?", "No. The selected v7 graph is fusion_dual_ulunas. GTCRN is retained as an older comparison architecture and project-history stage."],
    ["Does it use military data?", "Yes. MAD-labelled military sounds are indexed, weighted during training and tested in a dedicated military audit manifest."],
    ["Does it already prove full ANC?", "No. It proves communication speech enhancement/noise suppression. Physical acoustic cancellation needs a reference sensor, error microphone, actuator, secondary-path model and safe FxNLMS control loop."],
    ["Is Raspberry Pi 5 deployment proven?", "Not yet. Scripts and an execution plan exist; target-device latency, temperature, endurance and audio-I/O measurements are pending."],
])

doc.add_heading("1. Problem statement and acceptance interpretation", level=1)
doc.add_paragraph("SIH/DRDO Problem Statement 26052 asks for an AI/ML-enabled adaptive noise-cancellation system that suppresses stationary, non-stationary and impulsive defence noises while preserving speech intelligibility and operating in real time on embedded hardware. Its expected solution includes scalable pair generation, a robust model, optimized training, edge inference, microphones/headset integration, and target metrics of output SNR > 15 dB, STOI > 0.85 and PESQ > 2.5.")
table(doc, ["Requirement phrase", "VoxShield implementation", "Current status"], [
    ["Stationary noise", "MUSAN/ESC-50 and labelled mixture categories; long-duration backgrounds", "Implemented and evaluated"],
    ["Non-stationary noise", "Time-varying environmental, aircraft, rotor, vehicle and mixed events", "Implemented; hard military cases remain weak"],
    ["Impulsive noise", "Gunshot/artillery samples, transient-aware mixing and audit categories", "Implemented; still the broad-gate weak category in several stages"],
    ["Clean + noise at varied SNR", "Deterministic manifest generation with clean, near-clean and noisy conditions", "Implemented"],
    ["Full-band + sub-band", "UL-UNAS causal encoder/decoder with ERB-domain feature path", "Implemented"],
    ["Complex domain/phase", "Complex primary/reference features and complex ratio-mask output", "Implemented in v7"],
    ["Losses", "SI-SNR, SNR, MR-STFT, complex, waveform, identity, retention and artifact controls", "Implemented by configuration"],
    ["SNR/STOI/PESQ", "Metric runner, category/SNR breakdowns and strict gates", "Implemented; PESQ now computed in v6/v7 reports"],
    ["Real-time edge", "Stateful ONNX, streaming runtime, nonblocking worker and bypass", "Laptop verified; Pi pending"],
    ["Adaptive residual filtering", "NLMS experiments and separate FxNLMS path", "NLMS rejected for current v7 live chain; physical FxNLMS not yet proven"],
])

doc.add_heading("2. Scope: enhancement versus physical ANC", level=1)
doc.add_paragraph("The word ANC is often used loosely. VoxShield's completed neural path is a causal communication speech-enhancement system: it estimates cleaner speech that is delivered to headphones, a radio or a communication unit. It does not currently generate an anti-phase acoustic field that cancels sound at the user's ear. That second problem is physical active noise control and has different stability, geometry and hardware requirements.")
table(doc, ["Capability", "Current neural enhancement", "Physical ANC extension"], [
    ["Goal", "Make transmitted/listened speech clearer", "Reduce physical sound pressure at an error microphone"],
    ["Inputs", "Primary mic; optional noise-reference mic", "Reference mic + error mic + secondary-path measurement"],
    ["Output", "Enhanced audio samples", "Bounded anti-noise sent to actuator"],
    ["Core", "Causal complex-mask neural model", "FxNLMS/control law plus safety monitor; neural assistance optional"],
    ["Proof required", "Objective metrics, listening, streaming latency", "Closed-loop stability, cancellation dB at ear, actuator limits, fail-safe tests"],
])

doc.add_heading("3. Project evolution: what changed and why", level=1)
table(doc, ["Stage", "Major change", "Reason / result"], [
    ["Blank project", "Created repository, packaging, configs, data, models, training, evaluation, export, runtime, tests and docs", "No legacy code, data, checkpoints or test audio existed."],
    ["Early lightweight model", "~19.9k-parameter causal prototype", "Training loss improved, but evaluation showed negative SNR/SI-SNR improvement and audible noise retention. The model/curriculum was inadequate."],
    ["v3 official GTCRN", "Adopted official-style causal GTCRN, stronger losses and quality gates", "Strong gains on genuinely noisy samples, but clean/near-clean speech was degraded; aggregate results hid this regression."],
    ["v4 hybrid safeguards", "Added clean identity, speech retention, RIR/UAV data, label-aware audits and staged fine-tuning", "Improved intelligibility and preservation. Impulsive noise and clean-preservation trade-offs remained; stage2/2b gains were marginal."],
    ["v5 repair", "Acceptance workflow, stricter evidence, benchmark and regression diagnostics", "Exposed that a fast runtime does not mean acceptable enhancement quality."],
    ["v6 mono UL-UNAS", "Switched to official UL-UNAS family, pretrained DNS3 initialization and full-corpus fine-tuning", "Better modern mono baseline and the recommended fallback when only one valid microphone is available."],
    ["v7 dual fusion UL-UNAS", "Five ERB spatial features and dual-channel complex masking", "Current selected architecture; aggregate v7 gates are strongest, while the hard military tail remains below target."],
    ["Live engineering", "Nonblocking worker, health telemetry, safe aligned bypass, redesigned dashboard, duplicate-channel diagnostics", "Prevents the audio callback from waiting on inference and makes invalid laptop-array references visible."],
    ["Experimental dual-device", "Separate earphone mic + laptop array streams with acoustic delay calibration", "Makes a no-interface engineering demo possible, but clocks are unsynchronized and it is not final proof."],
])

doc.add_heading("4. Current repository and system modules", level=1)
table(doc, ["Area", "Responsibility", "Representative location"], [
    ["Configuration", "Architecture, optimizer, curriculum, losses, paths and gates", r"C:\anc\configs"],
    ["Dataset layer", "Recursive indexing, metadata parsing, splits and manifest generation", r"C:\anc\src\voxshield"],
    ["Training", "Mixed precision, logging, checkpointing, resume and best-model selection", r"C:\anc\src\voxshield\train.py"],
    ["Evaluation", "Objective metrics, breakdowns, strict acceptance gates and saved audio", r"C:\anc\src\voxshield\evaluate.py"],
    ["Runtime", "Stateful streaming inference, health modes, bypass and conditioning", r"C:\anc\src\voxshield\runtime.py"],
    ["Live I/O", "PortAudio capture/playback, worker coordination and session reports", r"C:\anc\src\voxshield\live.py"],
    ["Experimental dual I/O", "Two independent capture devices plus delay calibration", r"C:\anc\src\voxshield\live_dual_device.py"],
    ["Dashboard", "Meters, waveform, spectrum, waterfall and health diagnostics", r"C:\anc\src\voxshield\dashboard.py"],
    ["Deployment", "ONNX export, validation, benchmark and Raspberry Pi scripts", r"C:\anc\scripts"],
    ["Evidence", "JSON metrics, validation hashes, audio samples and session telemetry", r"C:\anc\artifacts"],
])

doc.add_heading("5. Dataset system", level=1)
doc.add_paragraph("The data pipeline reads source datasets in place; it does not require manually flattening large corpora. It recursively discovers supported audio formats, interprets available CSV/metadata and folder names, assigns a source/label/category, and writes lightweight indexes and mixture manifests. Deterministic group-aware splitting is used to reduce speaker/source leakage.")
table(doc, ["Dataset", "Role", "What it contributes"], [
    ["VCTK Corpus", "Clean target speech", "Multiple speakers, accents and recording variation"],
    ["MUSAN", "General noise", "Stationary/non-stationary background and music/noise diversity"],
    ["ESC-50", "Environmental events", "Siren, engine-like, natural and transient classes"],
    ["MAD_dataset", "Defence-domain noise", "Gunfire, artillery/shelling, tanks/vehicles, aircraft, helicopter/rotor and related military labels"],
    ["small_uav_acoustics", "UAV/rotor domain", "Recorded UAV flights and rotor signatures"],
    ["RIRS_NOISES", "Acoustic augmentation", "Room impulse responses and noise resources for reverberant propagation"],
])
callout(doc, "Indexed data snapshot", "The project status record reports 114,805 catalogued files, 39.33 hours of clean training speech, 1,573 gunshot files, 987 artillery/shelling files, 1,032 vehicle/tank files, 1,065 helicopter/rotor files, 1,031 aircraft files, 62 UAV flights and 60,000 RIR files. These are catalogue counts, not a claim that every file is unique defence audio or that every file appears in each training epoch.")
doc.add_heading("5.1 Expected external dataset locations", level=2)
add_bullets(doc, [
    r"C:\Users\jeswa\Downloads\Datasets\_n\VCTK-Corpus",
    r"C:\Users\jeswa\Downloads\Datasets\_n\musan",
    r"C:\Users\jeswa\Downloads\Datasets\_n\esc-50",
    r"C:\Users\jeswa\Downloads\Datasets\_n\MAD_dataset",
    r"C:\Users\jeswa\Downloads\Datasets\_n\small_uav_acoustics",
    r"C:\Users\jeswa\Downloads\Datasets\_n\RIRS_NOISES",
])
doc.add_heading("5.2 Mixture synthesis", level=2)
add_bullets(doc, [
    "Resample and condition speech/noise to the configured 16 kHz mono or paired-channel format.",
    "Select clean speech and one or more noises with deterministic seeds and group-aware split boundaries.",
    "Scale noise to the requested SNR, including clean, near-clean, normal, hard and very-hard conditions.",
    "Generate stationary, non-stationary, impulsive, environmental and mixed categories.",
    "Apply optional RIR convolution, reverberation, gain variation, clipping, spectral coloration and multiple-noise overlap.",
    "For v7 reference simulation, use inter-channel delay −16 to +16 samples, reference gain −6 to +6 dB, speech leakage −35 to −12 dB, reference self-noise SNR 35–55 dB and optional spectral tilt.",
    "Store noisy, clean, reference, labels, SNR and provenance in JSONL manifests so evaluation can be reproduced.",
])

doc.add_heading("6. Selected v7 architecture", level=1)
callout(doc, "Correct architecture name", "fusion_dual_ulunas: a causal dual-channel UL-UNAS-derived enhancement network. GTCRN is not fused into the selected v7 compute graph.", "EAF7F2", GREEN)
add_numbered(doc, [
    "Capture synchronized primary and noise-reference waveforms. The primary contains speech + noise; the reference should be noise-dominant with limited target-speech leakage.",
    "Apply safe input conditioning: finite-value checks, channel normalization/limits and framing at the configured 16 kHz rate.",
    "Compute a causal STFT representation and preserve recurrent/state caches between frames.",
    "Form five ERB-domain spatial features: primary log magnitude, reference log magnitude, cosine IPD, sine IPD and log-magnitude ratio.",
    "Feed the feature tensor to the widened causal UL-UNAS encoder/core/decoder.",
    "Estimate a two-channel complex ratio mask and apply it to the primary complex spectrum, preserving magnitude and phase relationships.",
    "Reconstruct samples with streaming inverse STFT/overlap-add, then use conservative wet/dry mixing, clipping protection and health-controlled bypass.",
])
table(doc, ["Property", "Current value"], [
    ["Architecture", "fusion_dual_ulunas"],
    ["Parameters", "172,748"],
    ["Sample rate", "16 kHz"],
    ["Streaming hop", "256 samples / 16 ms"],
    ["Inputs", "Primary + synchronized reference"],
    ["Output", "Enhanced single-channel speech"],
    ["Deployment", "Stateful ONNX opset 17"],
    ["Selected checkpoint", r"C:\anc\checkpoints_fusion_v7_stage2\best.pt"],
    ["Selected ONNX", r"C:\anc\artifacts\voxshield-v7-stage2-final.onnx"],
])

doc.add_heading("7. Training design", level=1)
doc.add_paragraph("Training optimizes a weighted multi-objective function rather than one metric. This is necessary because aggressive denoising can improve pause attenuation while making speech robotic, deleting consonants or changing active-speech level.")
table(doc, ["Loss component", "Purpose", "Failure it controls"], [
    ["SI-SNR", "Scale-invariant waveform fidelity", "Speech/noise separation quality"],
    ["Scale-sensitive SNR", "Absolute reconstruction fidelity", "Unwanted gain/level changes"],
    ["Multi-resolution STFT", "Magnitude consistency at several time-frequency resolutions", "Musical noise and spectral smearing"],
    ["Complex-domain loss", "Real/imaginary or complex-mask accuracy", "Phase distortion"],
    ["Waveform L1", "Sample-level regularization", "Transient and waveform mismatch"],
    ["Clean identity", "Output ≈ input on clean samples", "Clean speech damage"],
    ["Speech retention", "Penalize active speech deletion", "Dropped words/phonemes during impulses"],
    ["Spectral-hole / delta controls", "Protect continuity and spectral detail", "Robotic/tunnel voice"],
    ["Phase consistency", "Stabilize complex reconstruction", "Phasiness and metallic artifacts"],
])
doc.add_heading("7.1 Curriculum and checkpointing", level=2)
add_bullets(doc, [
    "Stage training rather than attempting one long blind run: broad learning, preservation tuning, then hard-military specialization.",
    "Use validation-controlled best checkpoints, periodic checkpoints, JSON logging, resume support and fresh-optimizer transfer when changing stage objectives.",
    "Include clean and near-clean samples in every later stage so military specialization does not erase speech-preservation behavior.",
    "Use source/label weights to increase hard but under-performing military categories without simply duplicating files.",
    "Stop or revise training when category gates regress; lower training loss alone is not acceptance evidence.",
])

doc.add_heading("8. Evaluation and acceptance methodology", level=1)
table(doc, ["Metric", "Interpretation", "Caution"], [
    ["SNR", "Energy ratio of clean target to reconstruction error", "Scale/alignment and silent regions affect it"],
    ["SI-SNR", "Scale-invariant waveform separation", "Can hide absolute level errors"],
    ["STOI", "Speech intelligibility estimate, roughly 0–1", "Objective proxy; not a listening replacement"],
    ["PESQ", "Perceptual speech quality estimate", "Narrow/wideband standard constraints; licensed implementation considerations"],
    ["Pause attenuation", "Noise reduction during speech pauses", "Can look excellent while active speech is harmed"],
    ["Speech deletion rate", "Fraction of active speech lost", "Thresholding must be interpreted with listening"],
    ["Active-level error", "Speech loudness preservation", "Does not alone measure timbre"],
    ["Latency/RTF", "Compute speed relative to frame duration", "Model time is not physical end-to-end audio latency"],
])
callout(doc, "Why aggregate and noisy-only numbers differ", "Clean and near-clean items are deliberately included to test preservation. Aggregate SNR improvement can be small or negative even while genuinely noisy items improve strongly. Therefore reports show both the complete set and a noisy-only subset, followed by category and SNR-bin gates.")

broad = load_json("artifacts/v7-stage2-onnx-broad-50.json")
mil = load_json("artifacts/v7-stage2-onnx-military-50.json")
bench = load_json("artifacts/v7-stage2-benchmark-onnx-cpu.json")
onnxv = load_json("artifacts/voxshield-v7-stage2-final.validation.json")

doc.add_heading("9. Verified results", level=1)
doc.add_heading("9.1 Current v7 sampled ONNX gates", level=2)
table(doc, ["Audit", "Output SNR", "Output STOI", "Output PESQ", "Meaning"], [
    ["Broad validation, 50 pairs", "26.87 dB", "0.902", "3.58", "Passes SIH absolute metric targets on this sampled aggregate gate"],
    ["Broad noisy-only subset", "17.17 dB", "0.867", "3.18", "Passes absolute targets on noisy-only broad subset"],
    ["Military validation, 50 pairs", "17.98 dB", "0.855", "2.86", "Passes absolute targets on sampled aggregate military gate"],
    ["Military hard/noisy-only subset", "10.66 dB", "0.816", "2.38", "Fails strict absolute targets; principal remaining quality gap"],
])
doc.add_paragraph("These values are evaluation results for fixed 50-pair audit manifests and the exported v7 ONNX artifact. They are not field-trial results, confidence intervals or proof of performance for every military environment.")
doc.add_heading("9.2 Hard military breakdown", level=2)
table(doc, ["Condition", "Output SNR", "STOI", "PESQ", "Assessment"], [
    ["Helicopter/rotor", "7.54 dB", "0.72", "1.90", "Major gap: periodic rotor harmonics overlap speech"],
    ["Fighter/aircraft", "8.89 dB", "0.76", "1.99", "Major gap: broadband/dynamic engine energy"],
    ["Vehicle/tank", "9.86 dB", "0.77", "2.13", "Below target: low-frequency masking and vibration-like spectra"],
    ["Input SNR < −10 dB", "5.25 dB", "0.74", "1.63", "Target speech often lacks recoverable evidence"],
    ["Input SNR ≥ 5 dB", "15.04 dB", "0.86", "2.93", "Meets the stated absolute targets in the reported bin"],
])
doc.add_heading("9.3 ONNX equivalence", level=2)
table(doc, ["Check", "Measured value", "Acceptance"], [
    ["Maximum enhanced-output absolute error", fmt(metric(onnxv, "max_output_absolute_error"), 10), "≤ 0.0002"],
    ["Maximum state absolute error", fmt(metric(onnxv, "max_state_absolute_error"), 10), "≤ 0.0002"],
    ["Result", str(metric(onnxv, "passed")), "Passed"],
    ["ONNX opset", str(metric(onnxv, "opset")), "17"],
])
doc.add_heading("9.4 Laptop CPU benchmark", level=2)
table(doc, ["Measure", "Result", "Interpretation"], [
    ["Mean ONNX inference", fmt(metric(bench, "mean_latency_ms"), 2) + " ms", "Below 16 ms frame duration"],
    ["p99 inference", fmt(metric(bench, "p99_latency_ms"), 2) + " ms", "Tail compute time"],
    ["Full-pipeline p99 callback", fmt(metric(bench, "full_pipeline_p99_callback_ms"), 2) + " ms", "Synthetic pipeline timing"],
    ["Model RTF", fmt(metric(bench, "real_time_factor"), 3), "< 1 means faster than real time"],
    ["Synthetic overruns", "2", "Requires margin and target-device measurement"],
])
callout(doc, "Latency claim boundary", "The neural hop is 16 ms and laptop compute is real time. This is not the same as microphone-to-headphone latency. Audio driver buffers, operating-system scheduling, resampling and device transport must be measured with a physical loopback.", "FFF1E8", ORANGE)

doc.add_heading("10. Historical results and lessons", level=1)
table(doc, ["Version", "Evidence", "What was learned"], [
    ["Early prototype", "Validation loss improved but first evaluation reduced SNR/SI-SNR", "Optimization loss alone was misleading."],
    ["v3 GTCRN", "Noisy-only test: +8.64 dB SNR, +7.71 dB SI-SNR, +0.054 STOI; full test clean/near-clean gate failed", "Strong denoising can coexist with speech damage."],
    ["v4 broad", "Noisy-only approximately +9.84 dB SNR, +9.56 dB SI-SNR, +0.113 STOI", "Intelligibility improved, but impulsive and preservation gates still exposed regressions."],
    ["v4 military", "Noisy-only approximately +13.60 dB SNR, +12.74 dB SI-SNR, +0.179 STOI", "Large gains on hard noise did not mean absolute military targets were met."],
    ["v6 UL-UNAS broad/full", "Output 21.32 dB SNR, 0.822 STOI, 2.83 PESQ; noisy-only +11.61 dB SNR", "Modern mono baseline improved quality and restored PESQ evaluation."],
    ["v6 UL-UNAS military/full", "Output 17.41 dB SNR, 0.775 STOI, 2.45 PESQ; noisy-only +12.07 dB SNR", "SNR target can pass while intelligibility/quality targets fail."],
    ["v7 fusion", "Aggregate sampled broad and military gates pass; hard military subset fails", "A real reference channel helps, but data severity and reference quality still dominate."],
])

doc.add_heading("11. Why robotic voice, wiped speech and distortion occurred", level=1)
table(doc, ["Observed symptom", "Technical cause", "Correction"], [
    ["Robotic/metallic voice", "Over-aggressive time-frequency mask, phase error, spectral holes, training/domain mismatch or wet mix too high", "Complex-domain supervision, artifact losses, clean identity, lower wet mix and listening gates"],
    ["Voice disappears during gunshot", "Impulse overlaps speech; low-SNR target evidence is masked and model mistakes target bins for noise", "Speech-retention loss, transient-over-speech curriculum, hard label weighting and a less aggressive wet/dry blend"],
    ["Noise remains clearly audible", "Early model capacity/features were insufficient or sample is outside training distribution", "UL-UNAS family, reference features, more domain data and category-aware fine-tuning"],
    ["Speaking triggers bypass", "Original duplicate-reference detector reacted after adaptive filtering or had unstable thresholds", "Detect physical similarity before LMS and add hysteresis"],
    ["Primary/reference meters almost equal", "Laptop microphone-array channels capture almost the same sound; measured correlation ≈0.99999 and normalized difference ≈0.0023", "Treat as invalid reference and use safe mono/bypass; use separated hardware for evidence"],
    ["Live audio breaks", "Inference inside callback, output underflows, driver selection, feedback or buffer mismatch", "Nonblocking worker, aligned dry fallback, health counters, headphones and controlled device APIs"],
    ["Telemetry PermissionError", "Multiple writers reused the same Windows .tmp path", "Unique temporary files, atomic replace retries and transient-write tolerance"],
])

doc.add_heading("12. Runtime safety and degraded modes", level=1)
add_bullets(doc, [
    "Audio callback never blocks waiting for inference; a worker processes queued frames.",
    "If output is late, the runtime uses aligned dry audio rather than stale or discontinuous enhanced audio.",
    "Invalid values, worker failure, excessive overruns or unusable reference conditions trigger degraded/bypass behavior.",
    "Wet mix provides conservative blending: output = wet × enhanced + (1 − wet) × aligned input.",
    "Raw-channel similarity is measured before any adaptive transformation; duplicated reference channels are rejected.",
    "Hysteresis prevents rapid normal/bypass flapping during speech.",
    "Telemetry records submitted/enhanced/bypass blocks, output underflows, dropped inputs, resets and failures.",
])
callout(doc, "Important decision", "Do not enable the evaluated NLMS pre-cancellation path in the current v7 live command. It caused severe speech regression. A physical FxNLMS loop is a separate hardware workstream, not a cosmetic switch.", "FDECEC", RED)

doc.add_heading("13. Engineering dashboard", level=1)
doc.add_paragraph("The dashboard is an observability interface, not the audio output. Audio is heard through the selected output device/headphones. The current UI shows runtime mode, inference time, wet mix, output reduction, primary/reference levels, channel correlation and normalized difference, underflows, waveform, spectrum, waterfall and health counters. It was redesigned because the earlier screen did not reveal whether the two channels were physically distinct or whether bypass was caused by invalid reference geometry.")
table(doc, ["Indicator", "Healthy interpretation", "Problem interpretation"], [
    ["Mode", "Normal", "Bypass/degraded means enhancement is intentionally limited"],
    ["Reference availability", "Present and meaningfully different", "N/A or duplicate indicates mono path/safe fallback"],
    ["Channel correlation", "Context dependent, not persistently ~1.0", "Near-perfect correlation suggests duplicate channels"],
    ["Normalized difference", "Clear nonzero spatial difference", "Near zero suggests same physical signal"],
    ["Underflows", "Stable at zero", "Increasing count predicts broken audio"],
    ["Wet mix", "Chosen for intelligibility/artifact trade-off", "1.0 can expose all model artifacts"],
])

doc.add_heading("14. Demonstration paths", level=1)
doc.add_heading("14.1 Recommended laptop mono demo", level=2)
doc.add_paragraph("Use the JBL/earphone microphone as a single primary microphone and headphones for output. Run the v6 mono artifact, which is the honest fallback when no independent reference microphone exists. Demonstrate raw versus enhanced recordings first, then live audio at a conservative wet mix around 0.55. Do not play the output through nearby speakers because feedback contaminates the result.")
doc.add_heading("14.2 Experimental laptop dual-device demo", level=2)
doc.add_paragraph("Use the earphone microphone as primary and the laptop array as a separate reference device. The experimental runtime opens independent streams, performs a short acoustic cross-correlation calibration and applies causal delay compensation. This can show the two-input concept without a capture interface, but independent sound-card clocks drift, calibration changes with movement and Windows scheduling can break synchrony. Label it an engineering demo, not the final tactical prototype.")
doc.add_heading("14.3 Proper dual-channel prototype", level=2)
add_bullets(doc, [
    "Use one synchronized two-input USB audio interface or codec.",
    "Place the primary mic near the talker's mouth and the reference closer to the dominant noise, with physical/acoustic shielding from speech where possible.",
    "Record real paired primary/reference/clean-proxy sessions for calibration and domain fine-tuning.",
    "Measure leakage, inter-channel delay, correlation, clipping and end-to-end latency.",
    "Use headphones/communication headset and add a hardware bypass/fail-safe.",
])

doc.add_heading("15. Raspberry Pi 5 deployment plan", level=1)
doc.add_paragraph("The Pi 5 target is feasible because the model is small and exported to ONNX, but feasibility must be demonstrated on the actual board. Raspberry Pi documents a 2.4 GHz quad-core 64-bit Arm Cortex-A76 CPU and recommends adequate power and active cooling for sustained workloads. The Pi must use a suitable USB audio device or HAT for the required microphone/headphone topology.")
add_numbered(doc, [
    "Install 64-bit Raspberry Pi OS, active cooling and a suitable power supply.",
    "Create a Python environment and install the project runtime/ONNX/audio dependencies using the Pi install scripts.",
    "Copy the selected ONNX artifact, configuration and any calibration data; verify hashes.",
    "List ALSA devices and first run offline inference on stored test files.",
    "Run a one-thread benchmark at 16 ms frames; record mean, p95, p99, RTF, memory, temperature and throttling.",
    "Run mono live audio with USB headset/interface. Confirm zero sustained underflows and safe bypass.",
    "For v7, use a synchronized dual-input interface and verify channel separation before enabling dual fusion.",
    "Measure physical microphone-to-headphone latency with loopback equipment; do not infer it from model timing.",
    "Run a two-hour soak test and capture telemetry, thermal behavior, failures and recovery.",
])

doc.add_heading("16. Feasibility, viability and deployment assessment", level=1)
table(doc, ["Dimension", "Assessment", "Evidence / condition"], [
    ["Technical feasibility", "Promising", "Causal 172,748-parameter stateful ONNX model; laptop real-time benchmark and objective gains"],
    ["Embedded feasibility", "Not yet proven", "Pi scripts exist, but target benchmark/thermal/soak evidence is pending"],
    ["Data feasibility", "Good for prototype", "Six data sources and reproducible manifest pipeline; real tactical paired recordings still needed"],
    ["Operational viability", "Conditional", "Needs rugged synchronized audio hardware, mic placement, calibration and fail-safe behavior"],
    ["Maintainability", "Good", "Config-driven modules, checkpoints, regression gates, tests and reproducible evidence artifacts"],
    ["Scalability", "Good", "Datasets are indexed in place; manifests and source weights avoid manual copying"],
    ["Privacy/security", "Requires deployment policy", "Prefer local inference; define retention, encryption, model update and audit controls"],
    ["Cost viability", "Potentially low at prototype scale", "Pi-class compute and compact model; final BOM depends on audio interface, headset, enclosure and qualification"],
])

doc.add_heading("17. Innovation and competition positioning", level=1)
add_bullets(doc, [
    "Defence-aware evaluation: separate military audits and label/SNR breakdowns, rather than reporting only a global average.",
    "Preservation-first acceptance: clean identity, speech deletion and active-level error explicitly prevent a false win from aggressive suppression.",
    "Dual-channel spatial fusion: primary/reference log magnitude, IPD and level-ratio information feed a compact causal neural core.",
    "Evidence-linked edge pipeline: the same checkpoint is exported, numerically validated, benchmarked and used by live runtime.",
    "Fail-operational design: aligned dry bypass and health telemetry are treated as system requirements, not UI extras.",
    "Clear separation of speech enhancement and physical ANC makes the roadmap technically credible.",
])
doc.add_paragraph("The strongest competition narrative is not that every metric is solved. It is that the team built a measurable engineering system, discovered where apparently good averages were misleading, corrected the architecture and acceptance process, and can show a concrete route from laptop evidence to a synchronized Pi prototype and, separately, to physical FxNLMS ANC.")

doc.add_heading("18. Risks and mitigations", level=1)
table(doc, ["Risk", "Impact", "Mitigation / evidence needed"], [
    ["Very low input SNR", "Speech information is physically masked", "Closer talker mic, directional/headset geometry, reference channel and realistic low-SNR curriculum"],
    ["Speech leakage into reference", "Model/adaptive filter may cancel target speech", "Placement, leakage measurement, training simulation and reference-quality gate"],
    ["Unsynchronized devices", "Phase/drift invalidates spatial features", "Single synchronized interface for final prototype"],
    ["Domain mismatch", "Field noises differ from public datasets", "Collect consented target hardware/room/vehicle recordings and fine-tune"],
    ["Robotic artifacts", "Lower perceived naturalness", "Conservative wet mix, perceptual listening, phase/spectral losses and clean gates"],
    ["Impulsive clipping", "Irrecoverable information loss", "Analog headroom, limiter, multi-level transient data and clipping augmentation"],
    ["Pi thermal throttling", "Latency spikes/underflows", "Active cooling, CPU governor tests, sustained soak benchmark"],
    ["Metric overclaim", "Judging credibility loss", "Show exact manifest size, hardware, artifact hash and aggregate/hard-subset distinction"],
])

doc.add_heading("19. Remaining engineering plan", level=1)
table(doc, ["Priority", "Work package", "Exit criterion"], [
    ["P0", "Run full, not only 50-pair, v7 broad and military evaluations", "Reproducible reports with category/SNR confidence and listening shortlist"],
    ["P0", "Listen to failure cases and audit manifests", "Confirmed labels, no corrupted/clipped targets, documented artifact taxonomy"],
    ["P0", "Run stage2c hard-military fine-tune", "Improved helicopter/fighter/vehicle and <−10 dB bins without clean/preservation regression"],
    ["P0", "Obtain synchronized dual-input hardware", "Distinct stable channels with measured leakage and delay"],
    ["P1", "Collect target-device paired recordings", "Field-like primary/reference data with metadata and consent"],
    ["P1", "Pi 5 benchmark and live validation", "RTF <1 with margin, no sustained underflows, thermal and two-hour soak report"],
    ["P1", "Physical loopback latency", "Measured microphone-to-headphone latency and jitter"],
    ["P2", "Physical FxNLMS prototype", "Secondary-path calibration, stable cancellation and hardware emergency bypass"],
    ["P2", "Deployment hardening", "Service supervision, signed artifact/version record, logs and recovery procedure"],
])
callout(doc, "Prepared next training configuration", r"C:\anc\configs\fusion_v7_stage2c_hard_military.yaml uses 8 epochs, learning rate 2×10⁻⁵, clean probability 0.25, near-clean probability 0.20, MAD source weight 5, and higher helicopter/fighter/vehicle weights. It should be evaluated against the same fixed audits to avoid moving the goalposts.")

doc.add_heading("20. Claims guide for the report and PPT", level=1)
table(doc, ["Safe to claim now", "Do not claim yet"], [
    ["Built end-to-end from scratch inside C:\\anc", "Field-proven defence deployment"],
    ["Six-source scalable dataset/indexing pipeline", "114,805 unique defence recordings"],
    ["Current v7 is a 172,748-parameter dual-channel fusion UL-UNAS", "Current v7 is UL-UNAS + GTCRN"],
    ["Sampled aggregate broad and military v7 gates meet stated SNR/STOI/PESQ targets", "All military categories and SNRs meet targets"],
    ["Hard military subset is the remaining gap", "Military problem is solved"],
    ["ONNX numerically matches PyTorch within tolerance", "INT8 is completed unless a validated quantized artifact exists"],
    ["Laptop CPU runs faster than real time in the reported benchmark", "Raspberry Pi latency is <30 ms without measuring it"],
    ["Safe bypass, telemetry and dashboard are implemented", "Bypass proves enhancement quality"],
    ["Experimental no-interface dual-device demo exists", "Separate consumer devices are synchronized tactical hardware"],
    ["Communication speech enhancement is implemented", "Acoustic anti-noise ANC is completed"],
])

doc.add_heading("21. Presentation-ready source content", level=1)
doc.add_heading("Core problem", level=2)
doc.add_paragraph("Defence communication is disrupted by overlapping stationary, dynamic and impulsive noise. Conventional stationary filters struggle when a gunshot, rotor harmonic, siren or vehicle spectrum changes faster than the filter can adapt. A useful system must suppress noise without deleting the very speech it is meant to protect, and it must do so causally on edge hardware.")
doc.add_heading("Solution statement", level=2)
doc.add_paragraph("VoxShield combines a reproducible defence-aware data pipeline, a compact causal dual-channel complex-mask network, objective and preservation-focused quality gates, stateful ONNX deployment, and a monitored real-time runtime with safe bypass. The present v7 prototype uses a primary microphone and a noise-reference microphone to exploit both spectral and spatial evidence.")
doc.add_heading("Key evidence", level=2)
add_bullets(doc, [
    "172,748 parameters; causal stateful ONNX export.",
    "ONNX-vs-PyTorch validation passed: enhanced-output maximum absolute error ≈1.04×10⁻⁷; state error ≈1.53×10⁻⁵.",
    "Sampled broad gate: 26.87 dB output SNR, 0.902 STOI, 3.58 PESQ.",
    "Sampled military aggregate gate: 17.98 dB output SNR, 0.855 STOI, 2.86 PESQ.",
    "Hard military subset: 10.66 dB, 0.816, 2.38—reported openly as the current gap.",
    "Laptop CPU benchmark: 6.91 ms mean inference and 11.41 ms p99 for a 16 ms frame; Pi measurement pending.",
])
doc.add_heading("Why it is credible", level=2)
doc.add_paragraph("Every major failure changed the design: negative SNR improvement led to better architecture/evaluation; clean-speech regression led to preservation losses and gates; robotic voice led to complex/phase and spectral artifact controls; duplicate laptop-array channels led to raw reference diagnostics and safe bypass; callback underflows led to nonblocking inference. This history demonstrates engineering iteration, not manufactured success.")
doc.add_heading("Demonstration sequence", level=2)
add_numbered(doc, [
    "Show a saved raw/noisy file and its enhanced counterpart with the same sample alignment.",
    "Show the evaluation JSON and explain aggregate versus hard/noisy-only results.",
    "Show ONNX equivalence and artifact hashes.",
    "Run the mono headset live demo for reliable audio; display latency/underflow telemetry.",
    "If distinct inputs are available, show the experimental dual-device concept with an explicit synchronization disclaimer.",
    "Show the Pi deployment checklist and identify target-device validation as the next milestone, not as completed evidence.",
])

doc.add_heading("22. Exact operational commands", level=1)
doc.add_paragraph("Run commands from PowerShell in C:\\anc. Device numbers can change after reconnecting hardware; list devices before every live demo.")
commands = [
    ("List audio devices", r'& ".\.venv\Scripts\voxshield-live.exe" --checkpoint ".\checkpoints_fusion_v7_stage2\best.pt" --onnx ".\artifacts\voxshield-v7-stage2-final.onnx" --list-devices'),
    ("Reliable mono laptop demo", r'& ".\scripts\run_v6_laptop_mono_demo.ps1" -InputDevice 8 -OutputDevice 12 -WetMix 0.55'),
    ("v7 laptop-array engineering demo", r'& ".\scripts\run_v7_laptop_demo.ps1" -InputDevice 9 -OutputDevice 12 -PrimaryChannel 0 -ReferenceChannel 1 -WetMix 0.70'),
    ("Experimental separate-device demo", r'& ".\scripts\run_v7_experimental_laptop_dual.ps1" -PrimaryDevice 8 -ReferenceDevice 9 -OutputDevice 12 -WetMix 0.70'),
    ("Prepared hard-military training", r'& ".\.venv\Scripts\voxshield-train.exe" --config ".\configs\fusion_v7_stage2c_hard_military.yaml" --device cuda'),
    ("Export selected model", r'& ".\.venv\Scripts\voxshield-export.exe" --checkpoint ".\checkpoints_fusion_v7_stage2\best.pt" --output ".\artifacts\voxshield-v7-stage2-final.onnx"'),
]
for title, command in commands:
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    r = p.add_run(command)
    r.font.name = "Cascadia Mono"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string("263238")
    p.paragraph_format.left_indent = Inches(0.2)

doc.add_heading("23. Verification checklist before submission", level=1)
add_bullets(doc, [
    "Freeze checkpoint, ONNX and manifest hashes; archive reports with date and hardware metadata.",
    "Run full broad and military sets, not only a 50-pair development sample.",
    "Install/verify PESQ in the exact evaluation environment and retain the metric implementation/version.",
    "Conduct blinded or at least structured listening on clean, near-clean, impulse-over-speech, rotor, vehicle and very-low-SNR samples.",
    "Demonstrate that clean speech is not harmed and speech-deletion rate remains below the stated bound.",
    "Benchmark Raspberry Pi 5 with one thread, active cooling and target audio I/O; retain thermal/soak telemetry.",
    "Measure real end-to-end latency using physical loopback.",
    "Use a synchronized dual-input device for the final dual-channel demonstration.",
    "Clearly label synthetic, public-dataset, laptop and physical-prototype evidence.",
])

doc.add_heading("24. Reference and standards trail", level=1)
refs = [
    ("UL-UNAS official repository", "https://github.com/Xiaobin-Rong/ul-unas"),
    ("UL-UNAS research paper", "https://arxiv.org/abs/2503.00340"),
    ("Ultra-low-latency speech enhancement research", "https://arxiv.org/abs/2409.10358"),
    ("ONNX Runtime documentation", "https://onnxruntime.ai/docs/"),
    ("ONNX Runtime IoT/edge guidance", "https://onnxruntime.ai/docs/tutorials/iot-edge/"),
    ("Raspberry Pi 5 official product page", "https://www.raspberrypi.com/products/raspberry-pi-5/"),
    ("Raspberry Pi 5 official product brief", "https://pip-assets.raspberrypi.com/categories/892-raspberry-pi-5/documents/RP-008348-DS-4-raspberry-pi-5-product-brief.pdf"),
    ("PortAudio API overview", "https://portaudio.com/docs/v19-doxydocs/api_overview.html"),
    ("PortAudio default-stream notes", "https://portaudio.com/docs/v19-doxydocs/open_default_stream.html"),
    ("Microsoft WASAPI documentation", "https://learn.microsoft.com/windows/win32/coreaudio/wasapi"),
    ("Microsoft loopback recording", "https://learn.microsoft.com/windows/win32/coreaudio/loopback-recording"),
]
for title, url in refs:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(title + ": ").bold = True
    p.add_run(url)

doc.add_heading("25. Final status statement", level=1)
doc.add_paragraph("VoxShield is beyond a concept: it has a complete data-to-deployment software chain, trained artifacts, objective reports, ONNX equivalence, live runtime, safety monitoring and a clear embedded plan. It is not yet a finished defence-qualified product. The honest completion boundary is: sampled aggregate targets achieved with v7; hardest military cases, synchronized hardware, Raspberry Pi validation, field data and physical ANC remain. That boundary should remain unchanged in the report, PPT and live judging narrative.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
