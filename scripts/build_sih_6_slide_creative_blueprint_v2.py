from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/VoxShield_SIH_6_Slide_Creative_Blueprint_V2.docx")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def style_table(table, header_fill: str = "0F172A") -> None:
    table.autofit = True
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            borders(cell)
            if r == 0:
                shade(cell, header_fill)
            elif r % 2 == 0:
                shade(cell, "F8FAFC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.3)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def h(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def p(doc: Document, text: str, lead: str | None = None):
    para = doc.add_paragraph()
    if lead:
        r = para.add_run(lead)
        r.bold = True
    para.add_run(text)
    return para


def bullets(doc: Document, items: list[str], numbered: bool = False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        doc.add_paragraph(item, style=style)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = head
    for values in rows:
        row = t.add_row()
        for i, value in enumerate(values):
            row.cells[i].text = value
    style_table(t)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t


def slide_section(
    doc: Document,
    n: str,
    title: str,
    hero: str,
    story: str,
    layout: str,
    content: list[str],
    visual: list[str],
    microcopy: list[str],
    notes: list[str],
):
    h(doc, f"Slide {n} {title}", 1)
    table(
        doc,
        ["Hero line", "Story job", "Layout idea"],
        [[hero, story, layout]],
        [2.1, 2.25, 2.9],
    )
    h(doc, "Final slide content", 2)
    bullets(doc, content)
    h(doc, "Visual build", 2)
    bullets(doc, visual)
    h(doc, "Small text and labels", 2)
    bullets(doc, microcopy)
    h(doc, "Speaker or video notes", 2)
    bullets(doc, notes)


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    for name in ("Normal", "List Bullet", "List Number"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.size = Pt(9.4)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("VoxShield Six Slide SIH Deck Blueprint V2")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("A complete creative, content, design, and video plan for Team@")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(70, 70, 70)

    p(
        doc,
        "This version is intentionally more detailed. It treats the six slide SIH limit as a design challenge. "
        "The deck should not feel like a report compressed into six pages. It should feel like a sharp defence technology briefing that a judge can understand almost instantly.",
    )

    h(doc, "The sharper concept", 1)
    p(
        doc,
        "The deck should sell one believable product story: Team@ built the AI engine and proof stack for a future field communication device. "
        "The live microphone problem should not control the narrative. The narrative should be controlled mission replay, local ONNX processing, edge benchmark, and hardware roadmap.",
    )
    table(
        doc,
        ["Layer", "What the judge sees", "What it silently proves"],
        [
            ["Mission pain", "A command is hidden by rotor, vehicle, gunfire, or siren noise.", "The problem is urgent and defence relevant."],
            ["System response", "VoxShield turns noisy speech into enhanced speech while showing telemetry.", "The solution is more than a model."],
            ["Engineering proof", "Data pipeline, ONNX export, Pi 5 timing, strict audit metrics.", "The prototype is measurable and reproducible."],
            ["Honest next phase", "Dual microphone hardware and field data roadmap.", "The team understands deployment reality."],
        ],
        [1.25, 3.0, 2.75],
    )

    h(doc, "The emotional hook", 1)
    p(
        doc,
        "The emotional hook should be a command that matters. Do not start with algorithms. Start with the moment where speech is lost.",
    )
    table(
        doc,
        ["Weak opening", "Stronger opening"],
        [
            ["Adaptive noise cancellation using AI ML", "A command should not disappear inside helicopter noise."],
            ["We use deep learning to remove noise", "VoxShield keeps speech readable when the environment turns hostile."],
            ["Our model improves SNR", "The operator sees the signal recover in real time and knows the frame deadline is safe."],
        ],
        [3.0, 3.8],
    )

    h(doc, "Deck visual language", 1)
    table(
        doc,
        ["Design element", "Use", "Why"],
        [
            ["Dark tactical base", "Deep navy, charcoal, black panels, low noise texture", "Feels like a mission console, not a classroom slide."],
            ["Signal green", "Passed checks, clean waveform, active model", "Creates instant positive status language."],
            ["Amber", "Known limitation, next phase, hardware risk", "Shows honesty without panic."],
            ["Waveform and spectrogram", "Before and after visuals on slides 1, 2, 5", "Audio becomes visible in seconds."],
            ["Metric badges", "ONNX passed, broad gate passed, Pi 5 real time", "Fast credibility anchors."],
            ["Few words", "Large nouns and short proof labels", "Works for 5 second judging."],
        ],
        [1.35, 3.0, 2.6],
    )

    slide_section(
        doc,
        "1",
        "Title page",
        "A command should not disappear inside defence noise",
        "Make the judge feel the mission problem before they evaluate the technology.",
        "Use the template title slide but make it cinematic. Left third: SIH metadata. Right two thirds: radio operator silhouette with noisy waveform becoming clean. Bottom strip: three proof badges.",
        [
            "VoxShield",
            "AI speech protection for noisy defence communication",
            "Problem Statement ID 26052",
            "Team@",
            "Edge ready prototype with mission replay, ONNX export, and Pi 5 benchmark",
        ],
        [
            "Background image idea: dark radio operator, headset silhouette, command post, or abstract battlefield communication line.",
            "Foreground graphic: one waveform starts amber/red and becomes signal green.",
            "Small badge 1: ONNX exported.",
            "Small badge 2: Broad strict gate passed.",
            "Small badge 3: Pi 5 real time benchmark.",
        ],
        [
            "PS Category: Hardware",
            "Theme: Miscellaneous",
            "Status: AI engine validated, hardware integration next",
        ],
        [
            "Video opening: play one second of noisy helicopter speech, then cut to this title.",
            "Do not explain the whole project here. Make it feel premium and serious.",
        ],
    )

    slide_section(
        doc,
        "2",
        "Idea title and proposed solution",
        "Noisy mission feed in, protected speech out",
        "Explain the solution through motion and system flow, not paragraphs.",
        "Center: VoxShield engine. Left: five noise sources. Right: two outputs, enhanced speech and mission telemetry. Bottom: short innovation ribbon.",
        [
            "Idea title: VoxShield Mission Communication Engine",
            "Proposed solution: local AI speech enhancement for defence noise with real time telemetry.",
            "The engine suppresses helicopter, vehicle, gunfire, siren, fan, and mixed noise while monitoring speech preservation.",
            "The dashboard gives operators bypass, safe, balanced, aggressive, mute, markers, and evidence export.",
            "The same ONNX model supports laptop replay and edge deployment testing.",
        ],
        [
            "Five noise source chips with icons: rotor, vehicle, gunfire, siren, fan.",
            "Central engine block with inner labels: dual channel deep filter, speech retention, ONNX runtime.",
            "Right output block: clear speech waveform plus dashboard screen.",
            "Innovation ribbon: model, telemetry, edge benchmark, failure analysis.",
        ],
        [
            "Local processing",
            "Speech deletion monitored",
            "Operator controlled",
            "Edge deployment path",
        ],
        [
            "This slide should feel like the product reveal.",
            "If a judge reads only one line, they should understand what VoxShield does.",
        ],
    )

    slide_section(
        doc,
        "3",
        "Technical approach",
        "The prototype covers data, model, runtime, and dashboard",
        "Show that the team built an end to end engineering stack.",
        "Use a left to right system architecture. Above the line show build pipeline. Below the line show runtime pipeline.",
        [
            "Build pipeline: datasets, mixture generator, training, hard failure repair.",
            "Runtime pipeline: noisy audio, reference channel, streaming model, enhanced speech, dashboard telemetry.",
            "Datasets: VCTK, DNS, MAD military audio, ESC 50, AudioSet targeted noise, RIRS, user recordings.",
            "Model: 1.1M parameter streaming deep filter dual channel network.",
            "Deployment: ONNX opset 17, parity checked, Pi 5 benchmark.",
        ],
        [
            "Diagram top row: data sources flow into training and hard failure report.",
            "Diagram bottom row: primary mic and reference mic flow into ONNX engine and output.",
            "Add a small inset showing frame processing: 16 kHz, 16 ms frame, stateful stream.",
            "Add one small dashboard screenshot or mockup thumbnail.",
        ],
        [
            "1.1M parameters",
            "ONNX parity passed",
            "16 ms frame deadline",
            "Pi 5 mean latency 5.33 ms",
        ],
        [
            "Narration: We did not stop at training. We built ingestion, evaluation, export, runtime, dashboard, and edge tests.",
            "Avoid showing too many model versions here.",
        ],
    )

    slide_section(
        doc,
        "4",
        "Feasibility and viability",
        "The AI engine is proven, the physical ANC hardware path is clear",
        "Turn honesty into credibility. Show what works and what still needs hardware.",
        "Top half: four proof cards. Bottom half: risk to mitigation roadmap. Use green for proven and amber for next phase.",
        [
            "Proof: v15 broad strict SIH gate passed.",
            "Proof: v12 remains the best metric benchmark candidate.",
            "Proof: ONNX export passed with parity below tolerance.",
            "Proof: Raspberry Pi 5 benchmark met real time timing.",
            "Risk: Bluetooth and laptop microphones do not provide reliable synchronized reference audio.",
            "Mitigation: controlled mission replay now, USB dual microphone hardware next.",
        ],
        [
            "Proof card 1: Broad v15 SNR 30.48, STOI 0.938, PESQ 3.94.",
            "Proof card 2: Pi 5 5.33 ms per frame, 16 ms deadline.",
            "Proof card 3: speech deletion 0.0 in audits.",
            "Proof card 4: ONNX export and parity passed.",
            "Roadmap: replay prototype, Pi 5 engine, dual mic rig, headset and radio integration, field fine tune.",
        ],
        [
            "Validated now: AI engine and edge inference.",
            "Next hardware phase: synchronized primary and reference microphones.",
            "Known hard case: vehicle and helicopter STOI reliability.",
        ],
        [
            "This slide should prevent technical judges from attacking the live Bluetooth issue.",
            "Say the limitation before they discover it. That makes the team look mature.",
        ],
    )

    slide_section(
        doc,
        "5",
        "Impact and benefits",
        "Clear commands can change outcomes in high noise environments",
        "Connect technical work to mission value and practical use.",
        "Use a before and after field communication scene. Left: command buried. Right: command readable with telemetry green.",
        [
            "Defence impact: clearer commands under helicopter, vehicle, gunfire, siren, and mixed noise.",
            "Operational impact: local processing reduces cloud dependence and network risk.",
            "Safety impact: speech deletion monitoring protects the speaker's voice.",
            "Engineering impact: dashboard shows when audio input, reference channel, or latency becomes unsafe.",
            "Scalability: defence radio, disaster response, aerospace, industrial safety, emergency services.",
        ],
        [
            "Before after transcript idea: Noisy input shows unreadable command fragments. Enhanced output shows readable command.",
            "Large spectrogram comparison: left dense noise, right speech bands more visible.",
            "Impact icons: soldier radio, rescue worker, aircraft crew, industrial headset.",
            "Use one emotional command example: Hold position or Move left.",
        ],
        [
            "Voice protected",
            "Local inference",
            "Operator aware",
            "Edge ready",
        ],
        [
            "This is the human slide. It should make the judge care.",
            "Keep one command example visible. Humans remember a sentence more than a metric.",
        ],
    )

    slide_section(
        doc,
        "6",
        "Research and references",
        "Evidence, sources, and next build phase",
        "End with proof and readiness, not a boring list of links.",
        "Make an evidence wall with four quadrants: data, metrics, artifacts, next hardware. Add QR placeholder if you have a demo video or GitHub later.",
        [
            "Data: VCTK, DNS Challenge, MAD military audio, ESC 50, AudioSet, RIRS NOISES, user recordings.",
            "Metrics: SNR, SI SNR, STOI, PESQ, speech deletion rate, active level error, latency.",
            "Artifacts: v12 best metric model, v15 broad strict pass, ONNX model, Pi 5 verification, failure report.",
            "Next build: synchronized dual microphone rig, raw USB audio interface, field recordings, headset integration.",
            "Final status: AI engine and edge feasibility validated. Physical ANC hardware validation is the next phase.",
        ],
        [
            "Use grouped logos or text chips for datasets, not long URLs.",
            "Use evidence badges: export passed, broad gate passed, Pi 5 real time, failure analysis done.",
            "Add mini timeline for next 30 days: dual mic rig, field recording, retrain, headset mockup.",
            "Keep detailed references in speaker notes if the portal allows notes in exported PDF.",
        ],
        [
            "Evidence ready",
            "Metrics reproducible",
            "Hardware path defined",
        ],
        [
            "This slide should leave the judge thinking: They know what they built and what remains.",
            "Avoid dumping every source link in tiny text.",
        ],
    )

    h(doc, "Slide density plan", 1)
    table(
        doc,
        ["Slide", "Maximum words", "Must have visual", "Must have proof"],
        [
            ["1", "35", "Hero communication image plus waveform", "ONNX, Pi 5, broad pass badges"],
            ["2", "70", "Noisy feed to VoxShield to clear speech pipeline", "Local ONNX processing label"],
            ["3", "85", "Two row architecture diagram", "1.1M parameters and 5.33 ms Pi 5 latency"],
            ["4", "90", "Readiness cards and hardware roadmap", "v15 broad pass and speech deletion 0.0"],
            ["5", "75", "Before after audio or transcript", "One metric card and one human command"],
            ["6", "85", "Evidence wall", "Dataset and artifact groups"],
        ],
        [0.75, 1.15, 3.1, 2.0],
    )

    h(doc, "Mission replay demo plan for video", 1)
    table(
        doc,
        ["Shot", "Duration", "Screen", "Voiceover"],
        [
            ["1", "0 to 8 s", "Noisy helicopter waveform fills screen. Command subtitle breaks.", "In mission communication, losing one command can break coordination."],
            ["2", "8 to 18 s", "VoxShield title and architecture flash.", "Team@ built VoxShield, a local AI speech enhancement engine for defence noise."],
            ["3", "18 to 45 s", "Mission console streams scenario. Toggle noisy and enhanced.", "For repeatable judging, this controlled mission replay runs the ONNX model locally frame by frame."],
            ["4", "45 to 65 s", "Latency and Pi 5 proof cards.", "The engine exports to ONNX and runs within real time frame deadlines on Raspberry Pi 5."],
            ["5", "65 to 90 s", "Evidence wall with metrics.", "Broad strict gate passes, speech deletion stays at zero in audits, and hard military cases guide the hardware phase."],
            ["6", "90 to 110 s", "Dual microphone hardware roadmap.", "The next step is synchronized primary and reference microphones for physical ANC validation."],
        ],
        [0.45, 0.75, 3.0, 3.0],
    )

    h(doc, "Visual asset checklist", 1)
    bullets(doc, [
        "One dark radio or headset hero image.",
        "One waveform before after graphic.",
        "One spectrogram before after graphic.",
        "One clean architecture diagram.",
        "One dashboard screenshot or high fidelity mockup.",
        "One Raspberry Pi 5 or edge device image.",
        "Noise category icons: helicopter, vehicle, gunfire, siren, fan.",
        "Evidence badges: ONNX passed, broad gate passed, Pi 5 real time, speech deletion monitored.",
    ])

    h(doc, "Final copy bank", 1)
    table(
        doc,
        ["Use case", "Line"],
        [
            ["One line pitch", "VoxShield protects voice commands in defence noise and runs locally on edge hardware."],
            ["Human hook", "A command should not disappear inside helicopter noise."],
            ["System claim", "The prototype links data generation, AI enhancement, ONNX deployment, edge timing, and operator telemetry."],
            ["Demo claim", "Controlled mission replay shows the model processing local audio frame by frame."],
            ["Honest limitation", "Physical ANC needs synchronized primary and reference microphones, which is the next hardware phase."],
            ["Selection close", "Team@ has a measurable AI engine and a clear path to a field communication device."],
        ],
        [1.5, 5.3],
    )

    h(doc, "What would make the deck feel premium", 1)
    bullets(doc, [
        "Use one strong visual per slide instead of many small decorations.",
        "Make metric cards look like instrument readouts, not normal table cells.",
        "Use consistent signal language: active, protected, verified, edge ready, hardware phase.",
        "Use short labels that feel operational: Mission feed, AI engine, Speech protected, Frame deadline, Reference health.",
        "Keep the limitation visually amber, not red. It is a roadmap item, not a project failure.",
        "Let the before after audio visual carry emotion. Do not overfill with explanation.",
    ])

    h(doc, "Emergency cut if the template space is too tight", 1)
    bullets(doc, [
        "Slide 1: keep only title, subtitle, ID, team, three badges.",
        "Slide 2: keep only solution pipeline and four innovation labels.",
        "Slide 3: keep architecture diagram and four metric tags.",
        "Slide 4: keep three proof cards and one hardware roadmap line.",
        "Slide 5: keep before after command example and four benefits.",
        "Slide 6: keep evidence wall and next hardware phase.",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
