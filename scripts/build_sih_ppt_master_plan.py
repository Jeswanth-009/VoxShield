from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/VoxShield_SIH_26052_PPT_and_Demo_Master_Plan.docx")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text_color(cell, color: str) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table, header_fill: str = "1F2937") -> None:
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                set_cell_text_color(cell, "FFFFFF")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F8FAFC")


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_bullets(doc: Document, items: list[str], style: str = "List Bullet") -> None:
    for item in items:
        p = doc.add_paragraph(style=style)
        p.add_run(item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = value
    style_table(table)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_slide_table(doc: Document, rows: list[list[str]]) -> None:
    add_table(
        doc,
        ["Slide", "Title", "Purpose", "Core content", "Visual direction"],
        rows,
        [0.45, 1.35, 1.65, 2.55, 1.6],
    )


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.size = Pt(12.5)
    styles["Heading 3"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("VoxShield SIH 26052 PPT and Demo Master Plan")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Team@  |  AI ML adaptive noise suppression for defence communication")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(60, 60, 60)

    doc.add_paragraph(
        "This document is the master source for building the final presentation, demo video, "
        "dashboard story, and judge-facing explanation for Problem Statement 26052. It is intentionally "
        "honest: VoxShield is currently a strong AI speech enhancement and edge inference prototype, not yet "
        "a fully proven physical headset ANC product. The winning strategy is to present it as a complete "
        "defence communication system prototype with evidence, deployment readiness, operator controls, and a "
        "clear hardware roadmap."
    )

    add_heading(doc, "Executive decision", 1)
    doc.add_paragraph(
        "The project should stop trying to win only by another small model improvement. Out of 500 teams, many "
        "will use AI. Team@ should compete on system completeness: data pipeline, model evolution, ONNX export, "
        "Pi 5 edge feasibility, mission dashboard, repeatable mission replay, operator controls, evidence pack, "
        "and honest hardware roadmap."
    )
    add_table(
        doc,
        ["Decision area", "Recommendation", "Reason"],
        [
            ["Final benchmark model", "Use v12 or v10 as metric-best validated candidate", "v12 has the strongest broad and military averages among the validated models, while v10 is the stable breakthrough baseline."],
            ["Latest experiment", "Keep v15 as real-world and hard-failure repair experiment", "v15 broad strict gate passes and averages are strong, but military STOI reliability still fails."],
            ["Demo strategy", "Use controlled real-time mission replay, not Bluetooth live mic", "Laptop and Bluetooth inputs do not provide synchronized primary and reference microphones and can make a good model look bad."],
            ["Claim wording", "AI speech enhancement engine with edge-ready deployment path", "This is true, defensible, and stronger than overclaiming physical ANC without proper hardware."],
        ],
        [1.35, 2.35, 3.35],
    )

    add_heading(doc, "Problem statement alignment", 1)
    add_table(
        doc,
        ["PS 26052 requirement", "How VoxShield addresses it", "Evidence or status"],
        [
            ["Stationary, non-stationary, and impulsive defence noise", "Training and audits cover fan, engine, vehicle, helicopter, siren, gunshot, shelling, footsteps, mixed noise, DNS, MAD, ESC-50, AudioSet, and user recordings.", "Manifest and audit pipeline exists; broad strict gate passes on v15."],
            ["Maintain speech intelligibility", "Training losses include SI-SNR, waveform, speech retention, spectral hole, clean identity, and STOI-oriented pressure.", "Broad STOI 0.938 on v15; military STOI average 0.915 but per-clip military STOI reliability remains a limitation."],
            ["Real-time performance on embedded hardware", "Model exported to ONNX and benchmarked on Raspberry Pi 5.", "Pi 5 mean latency about 5.33 ms per 16 ms frame; real-time factor about 0.33."],
            ["Prototype demonstration", "Mission replay console plus local processing evidence, with live telemetry and operator controls.", "Dashboard exists; next work is to upgrade replay mode and video flow."],
            ["Hardware pathway", "Primary and reference microphone geometry is identified as required for true physical ANC.", "Current laptop/Bluetooth demo is not proof of physical ANC; hardware phase is planned."],
        ],
        [1.65, 3.25, 2.25],
    )

    add_heading(doc, "What has been built so far", 1)
    add_bullets(doc, [
        "Dataset and manifest pipeline referencing more than 124000 audio files without copying audio.",
        "Curated sources include VCTK clean speech, MUSAN, ESC-50, MAD military audio, DNS Challenge, RIRS, targeted AudioSet clips, and user recordings.",
        "Dual-channel deep filtering model with ONNX export and streaming state support.",
        "Evaluation system with SNR, SI-SNR, STOI, PESQ, pause attenuation, speech deletion rate, level error, category checks, and strict SIH gates.",
        "Pi 5 deployment package and benchmark evidence showing real-time inference feasibility.",
        "Dashboard and live engine with telemetry, operator controls, waveform and spectrum panels.",
        "Failure analysis process that identifies remaining hard cases rather than hiding them.",
    ])

    add_heading(doc, "Model evolution story", 1)
    doc.add_paragraph(
        "The PPT should not present every experiment as a success. The story should show engineering maturity: "
        "baseline, breakthrough, failed attempts, data repair, and final evidence. This is more credible than a "
        "perfect linear story."
    )
    add_table(
        doc,
        ["Stage", "Meaning", "Outcome", "How to present"],
        [
            ["v7 and v8", "Early fusion and signed-noise baselines", "Broad gate passes but military performance weak.", "Use as baseline improvement starting point."],
            ["v10", "Major deep-filter breakthrough", "Strong averages, broad pass, military average strong, Pi deployable.", "Present as first serious validated model."],
            ["v11", "Real-world fine-tune attempt", "Regressed on metrics.", "Mention internally only if needed; not a main slide."],
            ["v12", "Hard-repair metric candidate", "Best overall metric model in saved comparisons.", "Use as final benchmark candidate."],
            ["v13", "STOI-hard repair attempt", "Did not improve military reliability.", "Do not promote."],
            ["v14", "User recordings and AudioSet data experiment", "Strong averages but broad strict failed.", "Shows data-first work and real-world direction."],
            ["v15", "Targeted hard-failure repair from v14", "Broad strict gate passed; military reliability still limited.", "Use as latest experimental evidence, not final product claim."],
        ],
        [0.8, 1.8, 2.35, 2.0],
    )

    add_heading(doc, "Core benchmark evidence to show", 1)
    add_table(
        doc,
        ["Model", "Broad gate", "Military gate", "Broad SNR STOI PESQ", "Military SNR STOI PESQ", "Use in story"],
        [
            ["v10", "Pass", "Fail reliability", "30.34  0.942  3.96", "24.63  0.917  3.66", "Stable breakthrough baseline."],
            ["v12", "Pass", "Fail reliability", "31.11  0.941  3.99", "25.40  0.919  3.69", "Best metric candidate."],
            ["v14", "Fail reliability", "Fail reliability", "29.15  0.930  3.86", "24.15  0.907  3.52", "Real-world data experiment."],
            ["v15", "Pass", "Fail reliability", "30.48  0.938  3.94", "25.15  0.915  3.63", "Latest hard-failure repair and demo candidate."],
        ],
        [0.7, 0.9, 1.05, 1.7, 1.8, 1.45],
    )
    doc.add_paragraph(
        "Important interpretation: averages are strong, but the military strict gate still fails because individual "
        "vehicle, helicopter, fighter, shelling, and mixed clips miss STOI reliability. This should be framed as a "
        "known deployment challenge, not hidden."
    )

    add_heading(doc, "Winning narrative for the PPT", 1)
    add_bullets(doc, [
        "Opening: Defence communication fails not only because of loud noise, but because noise changes rapidly and hides speech during critical commands.",
        "Gap: Classical ANC assumes stable noise and often fails under helicopter rotors, vehicles, gunfire, sirens, and mixed battle sounds.",
        "Insight: A useful solution cannot be only a model; it must be a deployable system with data generation, model export, edge inference, monitoring, and operator control.",
        "Solution: VoxShield is an AI speech enhancement engine packaged as an edge-ready mission communication prototype.",
        "Evidence: Show model evolution, broad strict pass, strong military averages, Pi 5 timing, speech deletion monitoring, and failure analysis.",
        "Demo: Controlled real-time mission replay, not unreliable Bluetooth mic live feed.",
        "Roadmap: Synchronized dual-mic hardware, field recordings, mono fallback, headset/radio integration.",
    ])

    add_heading(doc, "Slide by slide master deck", 1)
    add_slide_table(doc, [
        ["1", "Opening mission problem", "Create urgency", "One soldier communication line fails under helicopter, vehicle, gunfire and siren noise. State PS 26052 directly.", "Dark defence visual, waveform buried in noise, one sentence problem."],
        ["2", "Why ordinary ANC is not enough", "Show technical gap", "Stationary filters fail under non-stationary and impulsive noise. Bluetooth/laptop DSP is also unreliable.", "Simple comparison: classical ANC vs AI adaptive enhancement."],
        ["3", "Our answer VoxShield", "Introduce system", "AI/ML adaptive speech enhancement engine with dual-channel model, ONNX deployment, dashboard, and edge benchmark.", "System architecture hero diagram."],
        ["4", "What makes this more than a model", "Differentiate from other teams", "Dataset pipeline, model, evaluator, dashboard, Pi 5 deployment, evidence pack, failure analysis.", "Six-module product stack."],
        ["5", "Dataset and scenario coverage", "Prove realism effort", "VCTK, DNS, MAD, ESC-50, MUSAN, RIRS, AudioSet, user recordings. Noise categories: gunshot, vehicle, helicopter, siren, fan, mixed.", "Dataset map and category badges."],
        ["6", "Model evolution", "Show engineering process", "v7/v8 baseline, v10 breakthrough, v12 metric best, v14/v15 real-world repair.", "Timeline with honest pass/fail markers."],
        ["7", "Benchmark results", "Give proof", "Table with v10, v12, v15 broad and military metrics. Emphasize broad strict pass and strong military averages.", "Clean metric table with green and amber badges."],
        ["8", "Edge feasibility", "Show deployability", "Pi 5 benchmark: mean latency about 5.33 ms per 16 ms frame, real-time factor about 0.33.", "Latency gauge and edge device card."],
        ["9", "Mission replay demo setup", "Prepare judges", "Controlled local replay processes battlefield-like scenarios frame-by-frame using ONNX model.", "Dashboard screenshot/mockup and scenario selector."],
        ["10", "Demo scenario 1 helicopter", "Show impact", "Noisy vs enhanced speech; display waveform, spectrogram, latency, noise class.", "Before/after audio cards."],
        ["11", "Demo scenario 2 vehicle or gunfire", "Show robustness", "Show hard case improvement and explain limitations if needed.", "Split spectrum and speech protection meter."],
        ["12", "Operator console", "Make it product-like", "Safe, balanced, aggressive, bypass, mute, event markers, export report.", "Dashboard controls close-up."],
        ["13", "Honest limitations", "Build trust", "Not yet a proven physical ANC headset. Needs synchronized primary/reference microphones and field data.", "Limitation to solution roadmap, not apology."],
        ["14", "Hardware roadmap", "Show next phase", "USB dual-mic array, Pi 5 packaging, headset/radio integration, field fine-tune, mono fallback.", "Roadmap from current prototype to field kit."],
        ["15", "Why Team@ should be selected", "Close strongly", "Complete system prototype, measurable evidence, edge feasibility, transparent risks, fast path to hardware validation.", "Final product vision render and five differentiators."],
    ])

    add_heading(doc, "Dashboard and demo design specification", 1)
    add_table(
        doc,
        ["Panel", "Purpose", "Content to show"],
        [
            ["Mission Feed", "Make replay feel live", "Scrolling waveform and spectrogram for noisy, reference, and enhanced channels."],
            ["Scenario Selector", "Show defence relevance", "Helicopter, vehicle convoy, gunshot burst, siren, fan noise, mixed battle noise."],
            ["Operator Controls", "Make it system-like", "Safe, balanced, aggressive, bypass, mute, wet mix, gain, marker button."],
            ["Signal Health", "Explain why live hardware matters", "Channel correlation, reference leakage risk, clipping, input level, speech activity."],
            ["AI Engine", "Show local processing", "Model name, ONNX backend, checkpoint hash, frame size, sample rate, inference latency."],
            ["Speech Protection", "Address core concern", "Speech deletion rate, active level error, preservation badge, clean speech warning."],
            ["Noise Suppression", "Show benefit", "Before/after level reduction, band energy movement, category label, attenuation estimate."],
            ["Edge Readiness", "Differentiate", "Pi 5 latency, real-time factor, CPU use, memory, deployment status."],
            ["Evidence Export", "Credibility", "Save report, scenario ID, metrics, audio paths, timestamp."],
        ],
        [1.25, 1.7, 4.3],
    )

    add_heading(doc, "Video plan", 1)
    add_table(
        doc,
        ["Time", "Scene", "What viewer sees", "Narration"],
        [
            ["0 to 10 s", "Problem hook", "Noisy mission communication waveform and subtitles struggling to appear.", "In defence communication, the problem is not just loud noise. The problem is losing speech when noise changes suddenly."],
            ["10 to 25 s", "System overview", "Architecture: datasets to model to ONNX to dashboard to Pi 5.", "VoxShield is an AI speech enhancement prototype built as a deployable pipeline, not only a notebook model."],
            ["25 to 55 s", "Mission replay demo", "Dashboard streaming helicopter or vehicle scenario; toggle noisy/enhanced.", "For repeatable judging, we use controlled real-time replay. The model processes local audio frame by frame."],
            ["55 to 80 s", "Evidence view", "SNR, STOI, PESQ, speech deletion, latency and Pi 5 benchmark.", "We verify quality and runtime with objective metrics and edge timing."],
            ["80 to 105 s", "Hard case honesty", "Vehicle/helicopter hard clips highlighted.", "The current engine is strong, but physical ANC still needs synchronized microphone hardware and field data."],
            ["105 to 120 s", "Closing", "Roadmap and Team@ differentiators.", "Team@ is ready to convert this AI engine into a defence communication prototype with proper hardware validation."],
        ],
        [0.8, 1.35, 2.7, 2.4],
    )

    add_heading(doc, "PPT visual design guide", 1)
    add_bullets(doc, [
        "Theme: dark tactical interface, not colorful startup pitch. Use near-black background, muted grey panels, green signal accents, amber warnings, red only for failures.",
        "Fonts: Aptos or Inter for normal text; Cascadia Code or JetBrains Mono for telemetry numbers.",
        "Layout: one idea per slide, large visual center, short explanation on the side, metrics in compact cards.",
        "Color meaning: green means passed/real-time, amber means limitation/in progress, red means risk or failure only.",
        "Avoid: dense paragraphs, generic AI robot images, random military stock overload, unreadable charts, fake perfection.",
        "Use: waveform/spectrogram motifs, system diagrams, benchmark badges, scenario labels, hardware roadmap, evidence table.",
    ])

    add_heading(doc, "Exact positioning language", 1)
    add_table(
        doc,
        ["Use this wording", "Avoid this wording", "Why"],
        [
            ["AI ML adaptive speech enhancement engine", "Finished battlefield ANC headset", "The engine is validated, physical headset is not yet proven."],
            ["Controlled real-time mission replay", "Fake live demo", "Replay is legitimate if local model processing is real."],
            ["Edge-ready inference validated on Pi 5", "Fully deployed embedded product", "Benchmark proves feasibility, not packaged product completion."],
            ["Synchronized dual-mic hardware is the next phase", "Laptop mic proves ANC", "Laptop/Bluetooth does not represent correct ANC hardware."],
            ["Known limitation: military STOI reliability on vehicle/helicopter/fighter clips", "Model works perfectly everywhere", "Honesty increases judge trust."],
        ],
        [2.1, 2.1, 2.8],
    )

    add_heading(doc, "Best five demo examples to prepare", 1)
    add_bullets(doc, [
        "Helicopter plus speech: shows defence relevance and remaining hard-case work.",
        "Vehicle convoy plus speech: directly targets the current STOI weakness.",
        "Gunshot or shelling plus speech: shows impulsive noise handling.",
        "Siren or alarm plus speech: visually impressive in spectrum and easy for judges to understand.",
        "User-recorded fan or air noise: shows real local recording, even if not the primary proof.",
    ])

    add_heading(doc, "What judges are likely to ask", 1)
    add_table(
        doc,
        ["Question", "Best answer"],
        [
            ["Is this real time?", "The ONNX model runs in real time. On Raspberry Pi 5, mean latency was about 5.33 ms for a 16 ms frame, with zero simulated deadline overruns in the benchmark."],
            ["Is this true ANC?", "The current prototype validates the AI enhancement engine. True physical ANC requires synchronized primary and reference microphones; that is the planned hardware phase."],
            ["Why not just use noise suppression from Bluetooth headphones?", "Consumer Bluetooth DSP is closed, latency-heavy, and not controllable. VoxShield provides auditable model behavior, deployment metrics, and defence-noise training."],
            ["What is novel?", "The system combines defence-noise data generation, dual-channel AI enhancement, strict audit metrics, ONNX export, Pi 5 validation, dashboard telemetry, and failure analysis."],
            ["Where does it still fail?", "A small set of low-SNR vehicle, helicopter, fighter, and mixed military clips still miss STOI reliability. The mitigation is field data and synchronized hardware training."],
        ],
        [1.7, 5.5],
    )

    add_heading(doc, "Immediate execution checklist", 1)
    add_bullets(doc, [
        "Choose final benchmark model: v12 for metrics, v15 for latest demo if listening is better.",
        "Create mission replay dashboard mode using saved noisy/reference/enhanced files.",
        "Select five best audio examples and export noisy plus enhanced pairs.",
        "Record a two-minute video using dashboard replay, not Bluetooth mic live mode.",
        "Build PPT from the slide table in this document.",
        "Prepare one backup slide with honest limitations and hardware roadmap.",
        "Keep all artifact paths and metric JSONs ready for judge verification.",
    ])

    add_heading(doc, "Final story in one paragraph", 1)
    doc.add_paragraph(
        "Team@ built VoxShield for SIH Problem Statement 26052 as an AI ML adaptive noise suppression prototype for "
        "defence communication. The work evolved from early baseline models to a deployable deep-filter engine with "
        "ONNX export, strict audio quality audits, edge benchmark evidence on Raspberry Pi 5, and a mission-style "
        "operator dashboard. The best validated models show strong SNR, STOI, and PESQ improvements, broad strict "
        "gate passing, and zero speech deletion in audits. The project is not yet a finished physical ANC headset; "
        "the remaining step is synchronized dual-microphone hardware and field data collection. The selection case is "
        "that VoxShield is already a complete and measurable AI enhancement pipeline with a clear path to defence-grade "
        "hardware validation."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
