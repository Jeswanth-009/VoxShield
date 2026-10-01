from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/VoxShield_SIH_6_Slide_PPT_Manual_Build_Spec_V4.docx")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def style_table(table, header_fill="0B1F33") -> None:
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            borders(cell)
            if r == 0:
                shade(cell, header_fill)
            elif r % 2 == 0:
                shade(cell, "F8FAFC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.2)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def para(doc, text, lead=None):
    p = doc.add_paragraph()
    if lead:
        r = p.add_run(lead)
        r.bold = True
    p.add_run(text)
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
    for values in rows:
        row = t.add_row()
        for i, value in enumerate(values):
            row.cells[i].text = value
    style_table(t)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()


def slide_manual(
    doc,
    n,
    template_name,
    big_idea,
    reference,
    look,
    zone_rows,
    exact_text,
    build_steps,
    assets,
    mistakes,
):
    h(doc, f"Slide {n} {template_name}", 1)
    table(
        doc,
        ["Big idea", "Reference inspiration", "How it should feel"],
        [[big_idea, reference, look]],
        [2.25, 2.35, 2.35],
    )
    table(
        doc,
        ["Viewer scan order", "What they should understand"],
        [[row[0], f"{row[2]} — {row[3]}"] for row in zone_rows],
        [2.4, 4.4],
    )
    h(doc, "Canvas zones and exact placement", 2)
    table(doc, ["Zone", "Position", "Content", "Design treatment"], zone_rows, [0.9, 1.35, 3.25, 1.35])
    h(doc, "Exact text to paste", 2)
    bullets(doc, exact_text)
    h(doc, "Manual build steps in PowerPoint", 2)
    numbered(doc, build_steps)
    h(doc, "Visual assets to create or find", 2)
    bullets(doc, assets)
    h(doc, "Mistakes to avoid", 2)
    bullets(doc, mistakes)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    for name in ("Normal", "List Bullet", "List Number"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.size = Pt(9.3)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("VoxShield SIH 6 Slide PPT Manual Build Spec V4")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("A detailed do this exactly guide for manually building the SIH submission deck")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(70, 70, 70)

    para(
        doc,
        "This document is not another high level plan. It is a manual build specification for the six-slide SIH template. "
        "It tells you how each slide should look, where each visual should go, what text to paste, what diagrams to draw, and what to avoid. "
        "Use it while editing the PPT manually.",
    )

    h(doc, "Design decision after reviewing finalist references", 1)
    para(
        doc,
        "The finalist decks are dense. They do not win by looking empty or minimalist. They use packed but organized infographic pages with many proof objects. "
        "For VoxShield, the deck should look like a mission technology dossier, with dashboard visuals, audio signal graphics, metric cards, architecture blocks, and a hardware roadmap.",
    )
    table(
        doc,
        ["Reference pattern", "How to use it for VoxShield"],
        [
            ["Team Kingpin product render plus labelled mechanisms", "Use dashboard mockup, waveform panels, model engine block, and labelled hardware path."],
            ["CoalWorks dense architecture and stakeholder workflow", "Use a dense but organized system stack showing data, model, runtime, dashboard, and Pi 5."],
            ["Mindsmiths flow diagram and cost feasibility", "Use signal flow and real feasibility metrics instead of generic claims."],
            ["Telhan Sathi risk and mitigation slide", "Show Bluetooth/laptop hardware risk openly and give the dual-mic mitigation."],
            ["EarthBenders polished evidence wall and prototype section", "Make the final slide an evidence board with datasets, metrics, artifacts, QR/demo placeholder, and next phase."],
        ],
        [2.2, 4.7],
    )

    h(doc, "Global design system", 1)
    table(
        doc,
        ["Element", "Specification"],
        [
            ["Base style", "Use the SIH template structure, but fill it with tactical infographic content. Keep the official logos, footer, page number and required headings."],
            ["Color palette", "Deep navy #07111F, charcoal #111827, signal green #23D18B, amber #F59E0B, danger red #EF4444, white #FFFFFF, light grey #E5E7EB."],
            ["Fonts", "Use Aptos, Calibri or template default. Titles 28 to 34 pt. Section labels 13 to 16 pt. Body labels 9 to 12 pt. Metric numbers 18 to 26 pt."],
            ["Visual density", "Dense is allowed, but every dense block needs a title, border, icon and spacing. Do not let text float randomly."],
            ["Metric cards", "Dark rounded rectangle, green number, small white label, tiny source line. Use 3 to 4 per slide maximum."],
            ["Diagrams", "Use thick arrows, numbered blocks, and simple labels. No thin complex arrows everywhere."],
            ["Images", "Use one strong hero image or dashboard mockup per slide. Do not add random clipart just to fill space."],
            ["Tone", "Confident but honest. Say AI engine validated, physical ANC hardware next."],
        ],
        [1.4, 5.4],
    )

    h(doc, "Core facts allowed in the PPT", 1)
    bullets(doc, [
        "Problem Statement ID: 26052.",
        "Team Name: Team@.",
        "PS Category: Hardware.",
        "Core claim: AI speech enhancement engine for noisy defence communication.",
        "v15 broad strict SIH gate passed.",
        "v15 broad metrics: SNR 30.48 dB, STOI 0.938, PESQ 3.94.",
        "v15 military average metrics: SNR 25.15 dB, STOI 0.915, PESQ 3.63.",
        "Pi 5 benchmark: mean latency about 5.33 ms for a 16 ms audio frame.",
        "Model size: about 1.1M parameters.",
        "ONNX export and parity validation passed.",
        "Speech deletion rate: 0.0 in the reported audits.",
        "Known limitation: true physical ANC requires synchronized primary and reference microphones. Laptop/Bluetooth is not enough.",
    ])

    slide_manual(
        doc,
        "1",
        "Title page",
        "Make VoxShield look like a defence communication system before the judge reads details.",
        "Like Team Kingpin title page, but darker and more mission-oriented.",
        "A field communication boot screen: official SIH template plus one dramatic audio signal visual.",
        [
            ["1", "Top center", "SMART INDIA HACKATHON 2026 from template", "Keep original template styling."],
            ["2", "Left 45 percent", "Required metadata: PS ID, title, theme, category, team ID, team name", "Clean bullets. Do not overfill."],
            ["3", "Right 55 percent", "Hero visual: noisy waveform entering a VoxShield shield or dashboard frame", "Dark image, green audio line, soft glow."],
            ["4", "Bottom strip", "Three proof badges: ONNX passed, Pi 5 real-time, Broad strict gate passed", "Small rounded dark cards."],
        ],
        [
            "Title: VoxShield",
            "Subtitle: AI speech protection for noisy defence communication",
            "Problem Statement ID: 26052",
            "Team Name: Team@",
            "Badge 1: ONNX export passed",
            "Badge 2: Pi 5 real-time benchmark",
            "Badge 3: Broad strict gate passed",
        ],
        [
            "Keep the official SIH title layout and logos intact.",
            "Replace the center/right empty template space with a single hero composition.",
            "Draw or insert a waveform that starts jagged amber/red and becomes smoother green.",
            "Add a translucent dark rounded rectangle behind the proof badges.",
            "Keep all text readable after PDF export at 100 percent zoom.",
        ],
        [
            "Hero image prompt if generating: dark cinematic defence radio operator wearing headset, abstract waveform overlay, no weapons, no text, high contrast, navy and green lighting.",
            "Alternative if no image: create a dashboard-like black panel with waveform and grid lines.",
            "Icons for badges: chip, Raspberry Pi/edge device, check mark.",
        ],
        [
            "Do not add long PS description paragraph.",
            "Do not use a generic AI robot image.",
            "Do not make the title page white and empty like a normal school submission.",
        ],
    )

    slide_manual(
        doc,
        "2",
        "Idea title and proposed solution",
        "Explain the product in one glance: noisy mission audio goes in, protected speech and telemetry come out.",
        "Like Team Kingpin and EarthBenders proposed solution slides: problem, solution, uniqueness and prototype visual in one page.",
        "A product poster with a strong central engine block and surrounding proof modules.",
        [
            ["1", "Top band", "Idea title: VoxShield Mission Communication Engine", "Large bold title. Add subtitle below."],
            ["2", "Left column", "The problem: defence noises that hide speech", "Five noise chips with icons."],
            ["3", "Center", "VoxShield AI engine block", "Largest element on slide. Dark block with green edge."],
            ["4", "Right column", "Outputs: enhanced speech, dashboard telemetry, evidence report", "Three stacked cards."],
            ["5", "Bottom ribbon", "What makes it unique", "Four mini modules with icons."],
        ],
        [
            "Title: VoxShield Mission Communication Engine",
            "Subtitle: Local AI enhancement for clear commands under defence noise",
            "Problem chips: helicopter rotor, vehicle engine, gunfire, siren, fan and air noise",
            "Engine labels: dual channel deep filter, speech retention, ONNX runtime, live telemetry",
            "Output labels: enhanced speech, operator dashboard, evidence export",
            "Uniqueness ribbon: model plus dashboard, edge benchmark, strict audio audit, failure analysis",
        ],
        [
            "Draw a left-to-right flow: Noisy mission feed, VoxShield engine, protected speech output.",
            "Inside the engine block, put four small labelled layers. Do not write paragraphs.",
            "Put a small dashboard screenshot/mockup on the right card.",
            "Put noisy and enhanced mini-waveforms under the output.",
            "Use icons, not long sentences, for noise sources.",
        ],
        [
            "Noise icons: helicopter, truck, burst, siren, fan.",
            "Dashboard mockup: use current VoxShield console screenshot or create a simple black panel with waveform, spectrum and latency.",
            "Audio path diagram: thick green line through the center.",
        ],
        [
            "Do not say final battlefield headset.",
            "Do not write a paragraph explaining deep learning.",
            "Do not use vague phrases like next generation AI solution.",
        ],
    )

    slide_manual(
        doc,
        "3",
        "Technical approach",
        "Prove the engineering stack is real and complete.",
        "Like CoalWorks and EarthBenders technical slides: dense architecture map with many labelled subsystems.",
        "A two-layer architecture map: build pipeline on top, runtime pipeline below.",
        [
            ["1", "Top 20 percent", "Build pipeline: datasets, mixture generation, training, hard failure repair", "Numbered modules 1 to 4."],
            ["2", "Middle 45 percent", "Main architecture: primary/reference audio to ONNX model to enhanced output and dashboard", "Large central flow diagram."],
            ["3", "Right 25 percent", "Metric proof stack", "Four vertical metric cards."],
            ["4", "Bottom 10 percent", "Edge deployment line: laptop replay to Pi 5 to dual mic headset", "Small roadmap strip."],
        ],
        [
            "Data sources: VCTK, DNS, MAD, ESC-50, AudioSet, RIRS, user recordings",
            "Training: SNR variation, reverb, clipping, reference leakage, hard failure replay",
            "Model: streaming dual channel deep filter, 1.1M parameters",
            "Runtime: ONNX opset 17, 16 kHz audio, 16 ms frame, dashboard telemetry",
            "Edge: Pi 5 mean latency 5.33 ms per frame",
        ],
        [
            "Create a top row of small source blocks feeding into a training block.",
            "Create the main runtime row: Primary channel plus reference channel feed into VoxShield ONNX engine.",
            "From engine, split into enhanced speech and telemetry dashboard.",
            "On right, place metric cards: 1.1M params, ONNX parity passed, 5.33 ms frame, speech deletion 0.0.",
            "Use color coding: data blue, model green, telemetry purple, hardware amber.",
        ],
        [
            "Small dataset badges or logos if available.",
            "Model engine icon or chip icon.",
            "Pi 5 image or simple edge board icon.",
            "Dashboard thumbnail.",
        ],
        [
            "Do not list all model versions here.",
            "Do not make arrows so thin that they disappear in PDF.",
            "Do not use raw terminal screenshots.",
        ],
    )

    slide_manual(
        doc,
        "4",
        "Feasibility and viability",
        "Make the judge trust the project because you show proof and deployment risks honestly.",
        "Like Mindsmiths, Telhan Sathi and EarthBenders feasibility slides: proof cards plus challenge/mitigation.",
        "A readiness board: validated now, risk, mitigation, hardware path.",
        [
            ["1", "Top row", "Four proof cards", "Green check status."],
            ["2", "Middle left", "Feasibility meter", "Prototype readiness scale."],
            ["3", "Middle right", "Hardware path sketch", "Dual mic plus Pi 5 plus headset/radio."],
            ["4", "Bottom row", "Risk and mitigation table", "Amber labels, not red."],
        ],
        [
            "Proof card 1: v15 broad strict gate passed",
            "Proof card 2: ONNX export passed",
            "Proof card 3: Pi 5 real time benchmark, 5.33 ms per frame",
            "Proof card 4: speech deletion 0.0 in audits",
            "Risk: laptop and Bluetooth mics do not provide synchronized reference audio",
            "Mitigation: controlled mission replay now, USB dual mic hardware next",
            "Hard case: vehicle and helicopter STOI reliability",
        ],
        [
            "Make the top proof cards visually dominant.",
            "Show the hardware roadmap as five dots: AI engine, mission replay, Pi 5 runtime, dual mic rig, field headset validation.",
            "Add a small diagram of primary mic near mouth and reference mic facing noise.",
            "The risk table should be compact: Risk, why it matters, mitigation.",
        ],
        [
            "Green check icons.",
            "Amber warning icon.",
            "Small Pi 5 picture.",
            "Simple dual microphone geometry diagram.",
        ],
        [
            "Do not hide the live demo limitation.",
            "Do not make the risk card look like project failure.",
            "Do not claim physical ANC is already proven.",
        ],
    )

    slide_manual(
        doc,
        "5",
        "Impact and benefits",
        "Make the technical work matter to a human mission moment.",
        "Like finalist impact slides: audience map, benefits, and outcome flow.",
        "A before-after communication scene with audience benefits around it.",
        [
            ["1", "Center", "Command restoration visual", "Noisy transcript to enhanced transcript."],
            ["2", "Left", "Before: command hidden by noise", "Amber waveform/spectrogram."],
            ["3", "Right", "After: speech protected", "Green waveform/spectrogram."],
            ["4", "Bottom", "Impact audiences", "Four icons and benefit labels."],
        ],
        [
            "Command example: Hold position, left flank clear",
            "Before label: speech masked by helicopter and vehicle noise",
            "After label: speech readable after local enhancement",
            "Benefit 1: clearer commands in high noise",
            "Benefit 2: local processing without cloud dependence",
            "Benefit 3: operator sees signal health and latency",
            "Benefit 4: scalable to radios, headsets, rescue and industrial safety",
        ],
        [
            "Use a large before/after waveform or spectrogram pair.",
            "Place the command text over the waveform like a subtitle.",
            "Use icons for defence radio, rescue, aircraft/vehicle, industrial headset.",
            "Put one metric card: v15 broad SNR 30.48, STOI 0.938, PESQ 3.94.",
        ],
        [
            "Before waveform/spectrogram from actual noisy audio if possible.",
            "Enhanced waveform/spectrogram from v15 output if possible.",
            "Four audience icons.",
        ],
        [
            "Do not make this slide only a list of benefits.",
            "Do not use generic social/economic/environmental text unless tied to mission communication.",
            "Do not overpromise perfect noise removal.",
        ],
    )

    slide_manual(
        doc,
        "6",
        "Research and references",
        "End like an evidence board, not a bibliography.",
        "Like CoalWorks and EarthBenders final slides: research, market/evidence, links, demo proof.",
        "A four-quadrant evidence wall with a small next-phase roadmap and optional QR/demo link.",
        [
            ["1", "Top left", "Data sources", "Dataset chips."],
            ["2", "Top right", "Metrics and validation", "Metric chips."],
            ["3", "Bottom left", "Artifacts and evidence", "Files/reports list."],
            ["4", "Bottom right", "Next hardware phase", "Mini roadmap plus QR placeholder."],
        ],
        [
            "Data: VCTK, DNS Challenge, MAD military audio, ESC-50, AudioSet, RIRS NOISES, user recordings",
            "Metrics: SNR, SI-SNR, STOI, PESQ, speech deletion, active level error, latency",
            "Artifacts: v12 metric model, v15 broad strict pass, ONNX model, Pi 5 verification, failure report",
            "Next phase: synchronized dual microphone rig, raw USB audio interface, field recordings, headset/radio integration",
            "Final line: AI engine and edge feasibility validated, physical ANC hardware validation next",
        ],
        [
            "Use four bordered evidence boxes, not one long list.",
            "Use tiny link icons or QR placeholder for demo/video.",
            "Use short source names, not full long URLs.",
            "Put raw references in speaker notes or a very small footer if needed.",
        ],
        [
            "Dataset logos or simple text chips.",
            "QR code placeholder if demo video exists.",
            "Mini roadmap icon strip.",
            "Evidence folder/file icons.",
        ],
        [
            "Do not paste huge links across the slide.",
            "Do not make it look like a last-minute reference dump.",
            "Do not forget to delete slide 7 before final upload.",
        ],
    )

    h(doc, "Manual build checklist", 1)
    bullets(doc, [
        "Open the SIH PPT template and save a copy as VoxShield_TeamAt_SIH26052.pptx.",
        "Delete slide 7 only after using it as a rule reminder.",
        "Keep the required slide headings but improve the inside content.",
        "Build slide 3 first because it defines the system architecture visual style.",
        "Then build slide 2 using the same engine block from slide 3.",
        "Then build slide 4 proof cards using final numbers.",
        "Then build slide 5 before/after command restoration visual.",
        "Finally build slide 1 title and slide 6 evidence wall.",
        "Export to PDF and inspect at 100 percent zoom.",
    ])

    h(doc, "What the final deck should make judges feel", 1)
    bullets(doc, [
        "This is not another generic AI denoising idea.",
        "The team built a full data to deployment pipeline.",
        "The team understands real hardware constraints.",
        "The demo is controlled and believable.",
        "The project has enough proof to deserve selection for the next phase.",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
