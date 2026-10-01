from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/VoxShield_SIH_Reference_Inspired_6_Slide_Blueprint_V3.docx")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def style_table(t, header_fill="0B1F33") -> None:
    for r, row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            borders(cell)
            if r == 0:
                shade(cell, header_fill)
            elif r % 2 == 0:
                shade(cell, "F7FAFC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.4)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def para(doc, text, lead=None):
    p = doc.add_paragraph()
    if lead:
        r = p.add_run(lead)
        r.bold = True
    p.add_run(text)
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
    for row_values in rows:
        row = t.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = value
    style_table(t)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()


def slide(doc, n, template_title, reference_pattern, voxshield_version, layout, content, visuals, wow):
    h(doc, f"Slide {n} {template_title}", 1)
    table(
        doc,
        ["Reference pattern to borrow", "VoxShield adaptation", "5 second effect"],
        [[reference_pattern, voxshield_version, wow]],
        [2.4, 3.0, 1.6],
    )
    h(doc, "Layout", 2)
    bullets(doc, layout)
    h(doc, "Content to place", 2)
    bullets(doc, content)
    h(doc, "Visual assets and diagram treatment", 2)
    bullets(doc, visuals)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    for name in ("Normal", "List Bullet"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.size = Pt(9.4)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("VoxShield SIH Six Slide Deck Blueprint V3")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Reference inspired plan based on SIH finalist deck patterns")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(70, 70, 70)

    para(
        doc,
        "I inspected the provided SIH finalist and winner PDFs as visual references. The strongest decks use dense infographic storytelling, not plain minimal slides. They combine process diagrams, prototype images, proof badges, benefit maps, and concise labels. VoxShield should use that same high-density but organized style.",
    )

    h(doc, "What the reference decks teach us", 1)
    table(
        doc,
        ["Reference deck", "What works", "How VoxShield should use it"],
        [
            ["Team Kingpin", "Strong product rendering, orange solution blocks, technical slide packed with labelled parts, FEA/prototype proof.", "Use a strong mission console/dashboard visual and label the engine parts like a product system."],
            ["CoalWorks Arize", "Very dense but organized system architecture, many small modules, stakeholder flow, market and reference evidence board.", "Use a system-stack slide showing data, model, ONNX, dashboard, Pi 5, operator controls."],
            ["Mindsmiths", "Clear hardware flow diagram, cost/ROI feasibility, impact categories with icons and metrics.", "Use a signal flow diagram and feasibility proof cards with latency, ONNX, strict gate and hardware roadmap."],
            ["Telhan Sathi", "Workflow and implementation process are clear, feasibility uses risk/mitigation framing, impact is audience-centered.", "Use controlled mission replay workflow and honest laptop/Bluetooth limitation with mitigation."],
            ["EarthBenders", "Most polished infographic density, colored problem/solution cards, architecture map, prototype/future scope, QR/demo emphasis.", "Use an evidence wall, mission replay/demo QR, dashboard screenshots, and future hardware scope."],
        ],
        [1.3, 2.65, 3.0],
    )

    h(doc, "The new deck direction", 1)
    para(
        doc,
        "The deck should become a six-page mission dossier. Every slide should look useful even if the judge reads only the title and glances at the diagram. The design can be dense, but each dense area must be grouped into labelled modules. The winning feeling should be: this team built a system, tested it, knows the limitation, and can take it to hardware.",
    )
    table(
        doc,
        ["Old weak direction", "New reference inspired direction"],
        [
            ["Clean but empty slides", "Dense infographic slides with diagrams and proof tags"],
            ["Only model metrics", "Model metrics plus demo flow, dashboard, edge benchmark and hardware roadmap"],
            ["Generic AI solution", "Mission communication engine with operator console and evidence pack"],
            ["Hiding live demo weakness", "Controlled mission replay now, synchronized dual mic hardware next"],
            ["Long paragraphs", "Short labels, module cards, visual hierarchy, proof badges"],
        ],
        [3.0, 3.8],
    )

    h(doc, "Deck style to copy from the references", 1)
    bullets(doc, [
        "Use many visual elements, but group them. Reference decks are dense because every section has a label, border, icon, and visual anchor.",
        "Use one strong color identity. For VoxShield use tactical dark blue, signal green, amber, and white. Keep SIH blue footer if template requires it.",
        "Make slide 3 and slide 4 the densest. The finalists often use technical and feasibility slides to prove effort.",
        "Use proof objects. Prototype screenshot, dashboard mockup, architecture diagram, metric card, dataset badge, ONNX badge, Pi 5 badge.",
        "Avoid empty corporate white space. In SIH idea decks, empty slides look like lack of work.",
        "Use links/QR/demo proof on the last slide like winners do, but do not clutter with huge raw URLs.",
    ])

    slide(
        doc,
        "1",
        "Title page",
        "Reference decks keep the SIH title format but add a central graphic and brand identity.",
        "Make the title slide a mission system cover with a noisy to clean waveform and a dark radio communication image.",
        [
            "Left side: required SIH metadata in the template style.",
            "Center or right: VoxShield hero visual, headset/radio operator or mission dashboard silhouette.",
            "Bottom: three compact badges, ONNX passed, Pi 5 real time, Broad gate passed.",
            "Small subtitle under VoxShield: AI speech protection for noisy defence communication.",
        ],
        [
            "Problem Statement ID 26052.",
            "Team Name Team@.",
            "PS Category Hardware.",
            "VoxShield title in large type.",
            "One sentence: protects voice commands under helicopter, vehicle, gunfire and siren noise.",
        ],
        [
            "Use waveform line as a visual divider, red/amber on left and green on right.",
            "Use dark tactical image or abstract radio waves rather than generic AI brain.",
            "Keep the official SIH logo and template elements visible.",
        ],
        "Looks like a serious defence tech entry before any details.",
    )

    slide(
        doc,
        "2",
        "Idea title and proposed solution",
        "Team Kingpin and EarthBenders use this slide to show problem, solution, uniqueness and prototype in one crowded but readable poster.",
        "Create a one-slide product poster: noisy battlefield audio enters VoxShield, protected speech and mission telemetry exit.",
        [
            "Top: Idea title, VoxShield Mission Communication Engine.",
            "Left column: The problem, voice commands get masked by rotor, vehicle, gunfire, siren and mixed noise.",
            "Center: Large VoxShield engine block with inner labels, dual channel AI, speech retention, ONNX runtime, telemetry.",
            "Right column: Outputs, enhanced speech, operator dashboard, evidence report.",
            "Bottom strip: Why unique, model plus dashboard plus edge benchmark plus failure analysis.",
        ],
        [
            "Proposed solution: local AI speech enhancement engine with real time operator telemetry.",
            "Innovation: controlled mission replay, strict audio audit, ONNX export, Pi 5 edge proof.",
            "Honest scope: AI engine validated now, synchronized dual mic hardware next.",
        ],
        [
            "Use icon chips for helicopter, vehicle, gunfire, siren, fan.",
            "Use a fake but realistic mini dashboard image or screenshot as the right output.",
            "Use before/after waveform below the pipeline.",
        ],
        "Judge understands the whole idea in one glance.",
    )

    slide(
        doc,
        "3",
        "Technical approach",
        "CoalWorks and EarthBenders technical slides are dense architecture maps. They look impressive because they show many real subsystems.",
        "Make this the most detailed slide. Show VoxShield as a build pipeline and runtime pipeline.",
        [
            "Top band: Dataset layer, VCTK, DNS, MAD, ESC 50, AudioSet, RIRS, user recordings.",
            "Middle band: Training layer, mixture generation, SNR variation, reverb, clipping, reference leakage, hard failure repair.",
            "Core band: Model layer, streaming dual channel deep filter, 1.1M parameters, speech retention, STOI pressure.",
            "Runtime band: ONNX export, frame engine, dashboard telemetry, local processing.",
            "Edge band: Pi 5 benchmark, future dual mic headset/radio hardware.",
        ],
        [
            "Model: deep filter dual channel network.",
            "Audio: 16 kHz, 16 ms frame.",
            "Deployment: ONNX opset 17 with parity check.",
            "Benchmark: Pi 5 mean latency about 5.33 ms per frame.",
        ],
        [
            "Create a stacked architecture like EarthBenders slide 3.",
            "Use arrows and numbered stages, but make the stage labels bigger than the details.",
            "Add one small screenshot of dashboard or waveform plot.",
            "Add small badge cluster for ONNX, Pi 5, strict audit, speech deletion.",
        ],
        "Looks like a complete engineering system, not an ML notebook.",
    )

    slide(
        doc,
        "4",
        "Feasibility and viability",
        "Mindsmiths and Telhan Sathi use feasibility slides to split technical, economic, practical and risk mitigation. EarthBenders uses challenge versus solution cards.",
        "Show validated feasibility on top and hardware risks with mitigation below.",
        [
            "Top row proof cards: Broad strict gate passed, ONNX export passed, Pi 5 real time passed, speech deletion monitored.",
            "Center: readiness meter from prototype to field hardware.",
            "Bottom: risks and mitigation, laptop/Bluetooth reference issue, hard military STOI reliability, need synchronized microphones, field data collection.",
            "Right: small future hardware sketch, USB audio interface plus primary mic plus reference mic plus Pi 5.",
        ],
        [
            "v15 broad: SNR 30.48, STOI 0.938, PESQ 3.94.",
            "Military average: SNR 25.15, STOI 0.915, PESQ 3.63.",
            "Pi 5: 5.33 ms per frame, 16 ms frame deadline.",
            "Remaining issue: vehicle and helicopter STOI reliability.",
        ],
        [
            "Use green proof cards and amber risk cards.",
            "Use a small hardware chain diagram, not only text.",
            "Use a maturity scale: AI engine, replay demo, edge benchmark, dual mic hardware, field validation.",
        ],
        "Judges see that you know what is proven and what comes next.",
    )

    slide(
        doc,
        "5",
        "Impact and benefits",
        "Finalist impact slides are audience maps, benefit wheels, or journey flows. They show who benefits and how.",
        "Make this slide about command clarity and operational trust.",
        [
            "Center: one command restored, noisy transcript unreadable, enhanced transcript readable.",
            "Around it: four impact groups, soldiers, command operators, disaster response, industrial safety.",
            "Benefit cards: clearer commands, local inference, speech protection, operator awareness, scalable to headset/radio.",
            "Bottom: before after waveform or spectrogram.",
        ],
        [
            "Defence: supports mission communication in rotor, vehicle, gunfire, siren and mixed noise.",
            "Operational: local processing reduces cloud dependency.",
            "Safety: speech deletion monitoring prevents voice removal.",
            "Economic and deployment: low size model, Pi 5 feasibility, path to portable hardware.",
        ],
        [
            "Use one human command example, Hold position or Move left.",
            "Use a circular impact map like Mindsmiths or Telhan Sathi, but with VoxShield categories.",
            "Use green enhanced transcript and amber noisy transcript.",
        ],
        "Turns the technical work into human mission value.",
    )

    slide(
        doc,
        "6",
        "Research and references",
        "CoalWorks and EarthBenders make the final slide an evidence board, not a plain bibliography.",
        "Create a research and evidence wall with datasets, metrics, artifacts, demo links and next phase.",
        [
            "Quadrant 1: Data sources, VCTK, DNS, MAD, ESC 50, AudioSet, RIRS, user recordings.",
            "Quadrant 2: Metrics, SNR, SI SNR, STOI, PESQ, speech deletion, latency.",
            "Quadrant 3: Artifacts, v12 metric candidate, v15 broad pass, ONNX model, Pi 5 verification, failure report.",
            "Quadrant 4: Next phase, synchronized dual mic rig, field recordings, headset/radio integration.",
            "Optional bottom right: QR code to demo video or local evidence folder if available.",
        ],
        [
            "Use grouped sources instead of long URLs.",
            "List 4 to 6 references only, with short names.",
            "Add status line: AI engine validated, physical ANC hardware validation next.",
        ],
        [
            "Use evidence cards and small QR placeholder.",
            "Use mini timeline from prototype to hardware validation.",
            "Put raw URLs in notes or tiny reference strip only if required.",
        ],
        "Ends with credibility and readiness instead of a boring reference dump.",
    )

    h(doc, "What VoxShield should copy and avoid from the references", 1)
    table(
        doc,
        ["Copy this", "Avoid this"],
        [
            ["Dense architecture slides with labelled modules", "Tiny unreadable text blocks with no grouping"],
            ["Prototype and dashboard visuals", "Only abstract icons"],
            ["Risk and mitigation framing", "Pretending no limitation exists"],
            ["Benefit maps connected to real users", "Generic social economic environmental filler"],
            ["Evidence board final slide", "Raw URL list without meaning"],
            ["Proof badges and metrics", "Too many metrics with no story"],
        ],
        [3.3, 3.3],
    )

    h(doc, "Final six slide storyboard", 1)
    table(
        doc,
        ["Slide", "Judge should think", "Main visual"],
        [
            ["1", "This is a serious defence communication idea.", "Radio operator plus waveform recovery."],
            ["2", "I understand the product in one glance.", "Noisy feed to VoxShield to clear speech and telemetry."],
            ["3", "They built a real system stack.", "Layered architecture map."],
            ["4", "They have proof and understand deployment risk.", "Proof cards plus hardware roadmap."],
            ["5", "This matters in real missions.", "Command restored before after visual."],
            ["6", "They have evidence and a next plan.", "Research and artifact evidence wall."],
        ],
        [0.65, 3.2, 3.0],
    )

    h(doc, "Immediate action list to build the actual PPT", 1)
    bullets(doc, [
        "Use the SIH template and delete slide 7.",
        "Create or capture one dashboard screenshot or mockup for slides 2, 3 and 5.",
        "Generate a noisy to enhanced waveform graphic for slides 1, 2 and 5.",
        "Create one architecture diagram for slide 3 using grouped blocks.",
        "Create four proof cards for slide 4 with v15 broad, military average, ONNX parity and Pi 5 latency.",
        "Create a QR code only if you have a demo video link ready.",
        "Export as PDF because SIH portal accepts PDF, not PPTX.",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
