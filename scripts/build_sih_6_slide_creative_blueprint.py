from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/VoxShield_SIH_6_Slide_Creative_Blueprint.docx")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "6")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)


def style_table(table, header_fill: str = "111827") -> None:
    table.autofit = True
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            borders(cell)
            if r == 0:
                shade(cell, header_fill)
            elif r % 2 == 0:
                shade(cell, "F8FAFC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(8.5)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def h(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def para(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def bullets(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = head
    for values in rows:
        row = t.add_row()
        for i, value in enumerate(values):
            row.cells[i].text = value
    style_table(t)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return t


def slide_block(
    doc: Document,
    number: str,
    title: str,
    emotional_job: str,
    five_second_read: str,
    layout: str,
    exact_copy: list[str],
    visuals: list[str],
    human_touch: str,
    avoid: str,
):
    h(doc, f"Slide {number} {title}", 1)
    table(
        doc,
        ["Creative job", "Five second read", "Composition"],
        [[emotional_job, five_second_read, layout]],
        [2.2, 2.0, 3.0],
    )
    h(doc, "Exact slide copy", 2)
    bullets(doc, exact_copy)
    h(doc, "Visual treatment", 2)
    bullets(doc, visuals)
    para(doc, human_touch, "Human feel: ")
    para(doc, avoid, "Avoid: ")


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    for name in ("Normal", "List Bullet"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.size = Pt(9.5)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 2"].font.size = Pt(12)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("VoxShield Six Slide SIH Creative Blueprint")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("A judge first deck plan for Team@ and Problem Statement 26052")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(65, 65, 65)

    para(
        doc,
        "This document rewrites the six slide plan as a creative judging deck. The goal is not to dump all work into the template. "
        "The goal is to make a judge understand, in a few seconds, that Team@ has built a serious defence communication prototype with evidence, local processing, edge feasibility, and a realistic hardware path.",
    )

    h(doc, "The creative idea", 1)
    para(
        doc,
        "Use the deck as a mission briefing. Each slide should feel like a page from a field system dossier, not a college project report. "
        "The visual language should say signal, pressure, field noise, and command clarity. Use fewer claims and more proof objects: waveform, spectrogram, latency, model hash, Pi benchmark, audio scenario, and dashboard.",
    )
    table(
        doc,
        ["Weak deck feeling", "Winning deck feeling"],
        [
            ["We trained an AI model", "We built a field communication engine"],
            ["Paragraphs explaining ANC", "Mission visuals showing speech rescued from noise"],
            ["Many metrics with no emotion", "One strong number beside a before after signal"],
            ["Live mic struggle", "Controlled mission replay with real local ONNX processing"],
            ["Future scope as excuse", "Hardware roadmap as deployment plan"],
        ],
        [3.1, 3.7],
    )

    h(doc, "One line that should drive the whole deck", 1)
    para(
        doc,
        "VoxShield protects voice commands in defence noise and proves the AI engine can run locally on edge hardware.",
    )
    para(
        doc,
        "This is the sentence every slide should support. If a line, image, or metric does not support it, remove it.",
    )

    h(doc, "Judge scan strategy", 1)
    table(
        doc,
        ["What judges do", "What the slide must do"],
        [
            ["They skim fast", "Use a large visual anchor and one dominant sentence."],
            ["They compare hundreds of AI ideas", "Show system completeness, not only model accuracy."],
            ["They distrust overclaims", "State the physical hardware limitation clearly and turn it into the next phase."],
            ["They remember demos", "Make mission replay look like a controlled field test with local processing."],
            ["They look for feasibility", "Show ONNX export, Pi 5 latency, dashboard, and evidence files."],
        ],
        [2.0, 4.8],
    )

    slide_block(
        doc,
        "1",
        "Title page",
        "Make the judge feel the seriousness of a defence communication problem before they read anything else.",
        "Team@ is solving command clarity under battlefield noise.",
        "Full bleed dark field communication image. Left side has title and SIH metadata. Right side has one waveform that starts chaotic and becomes clean. Bottom has three small proof badges: ONNX, Pi 5, Mission Console.",
        [
            "VoxShield",
            "AI speech protection for noisy defence communication",
            "Problem Statement ID 26052",
            "Team@",
            "Edge ready prototype with mission replay and Pi 5 benchmark",
        ],
        [
            "Background: soldier radio silhouette, headset, command post, or abstract tactical audio map.",
            "Overlay: red or amber noisy waveform fading into green enhanced waveform.",
            "Small badges: ONNX exported, Pi 5 real time, Speech deletion monitored.",
            "Keep SIH required metadata clean and aligned. Do not make the title slide crowded.",
        ],
        "The slide should feel like a system boot screen, not a school cover page. Make the title short and confident.",
        "Do not put dataset names, long PS title paragraphs, or all metrics on the title slide.",
    )

    slide_block(
        doc,
        "2",
        "Idea title and proposed solution",
        "Show the full product idea in one picture. The judge should understand the solution without reading a paragraph.",
        "Noisy mission audio enters VoxShield and clear speech with telemetry comes out.",
        "Use a large central pipeline shaped like an audio signal path. Left has field noises. Middle has VoxShield AI engine. Right has clear speech and operator dashboard.",
        [
            "Idea title: VoxShield Mission Communication Engine",
            "Proposed solution: local AI speech enhancement for defence noise, with operator telemetry and edge deployment readiness.",
            "Input: helicopter, vehicle, gunfire, siren, fan, mixed noise.",
            "Engine: dual channel deep filter model, speech retention loss, ONNX runtime.",
            "Output: enhanced speech, latency telemetry, evidence export.",
        ],
        [
            "Use icons or image chips for helicopter, vehicle, gunshot, siren, and fan.",
            "Use one central black box labelled VoxShield AI Engine.",
            "Show two output branches: Clear speech and Mission console.",
            "Add a tiny waveform before and after under the pipeline.",
        ],
        "This slide should feel like you are revealing a product, not explaining a model. The phrase mission communication engine gives it more weight.",
        "Do not use generic AI terms like revolutionary, next generation, or cutting edge. Let the system diagram create confidence.",
    )

    slide_block(
        doc,
        "3",
        "Technical approach",
        "Prove that serious engineering work sits behind the idea.",
        "Data to model to ONNX to edge console is already built.",
        "Split slide into a vertical build stack on left and a large architecture diagram on right. The architecture should be readable in three seconds.",
        [
            "Data layer: VCTK, DNS, MAD, ESC 50, AudioSet, RIRS, user recordings.",
            "Training layer: SNR variation, reverb, clipping, reference leakage, hard failure replay.",
            "Model layer: streaming dual channel deep filter, 1.1M parameters.",
            "Runtime layer: ONNX opset 17, parity checked, CPU inference.",
            "Interface layer: dashboard telemetry, operator controls, mission replay.",
        ],
        [
            "Main visual: five layer architecture stack with a signal line flowing through it.",
            "Add a small metric tag beside model: 1.1M parameters.",
            "Add a small export tag: ONNX parity passed.",
            "Add a small edge tag: 5.33 ms per 16 ms frame on Pi 5.",
        ],
        "Make it feel like a defence system pipeline. Data, model, runtime, dashboard, and hardware path should appear as one connected machine.",
        "Do not show code screenshots. Do not list every script. Do not make a dense table of all versions.",
    )

    slide_block(
        doc,
        "4",
        "Feasibility and viability",
        "Make the judge trust you because you are honest and evidence driven.",
        "The engine works locally, the edge path is proven, and the hardware gap is understood.",
        "Use three evidence cards across the top and a risk to mitigation strip at the bottom. The slide should look like an engineering readiness report.",
        [
            "Evidence 1: ONNX export passed with parity below tolerance.",
            "Evidence 2: Broad strict gate passed on v15.",
            "Evidence 3: Pi 5 benchmark passed with 5.33 ms mean latency per 16 ms frame.",
            "Known risk: laptop and Bluetooth microphones do not provide synchronized reference audio.",
            "Mitigation: controlled mission replay now, USB dual microphone hardware next.",
        ],
        [
            "Use green evidence cards for passed items.",
            "Use one amber card for the honest hardware limitation.",
            "Show a tiny hardware roadmap: laptop replay, Pi 5 engine, USB dual mic, headset radio integration.",
            "Add a small badge: Speech deletion 0.0 in audits.",
        ],
        "This is where honesty becomes a strength. Say what is proven and what is next. Judges prefer a team that understands deployment risk.",
        "Do not pretend Bluetooth live mic proves physical ANC. That can backfire if they ask technical questions.",
    )

    slide_block(
        doc,
        "5",
        "Impact and benefits",
        "Move from engineering proof to human value. Show why command clarity matters.",
        "VoxShield can protect mission speech where noise hides critical commands.",
        "Use a large before after communication scene. Left side: command lost in noise. Right side: speech restored and operator console confirms safe output.",
        [
            "Primary impact: clearer commands in helicopter, vehicle, gunfire, siren, and mixed-noise environments.",
            "Operational benefit: local processing avoids cloud dependence.",
            "Safety benefit: speech deletion monitoring protects the speaker's voice.",
            "Deployment benefit: edge benchmark supports portable field devices.",
            "Expansion: defence radio, disaster response, aerospace, industrial safety.",
        ],
        [
            "Use a before after audio strip: noisy transcript unreadable, enhanced transcript readable.",
            "Show waveform or spectrogram before after, but keep it large and visual.",
            "Use four impact icons: soldier radio, edge chip, safety shield, industrial headset.",
            "Put one metric card in the corner: v15 broad SNR 30.48, STOI 0.938, PESQ 3.94.",
        ],
        "Add a small human moment: a single command such as hold position or move left becoming readable after enhancement. That creates emotional memory.",
        "Do not make benefits abstract. Avoid social economic environmental filler unless the template forces it. Keep it mission focused.",
    )

    slide_block(
        doc,
        "6",
        "Research and references",
        "End with credibility. Show that claims connect to data, metrics, and reproducible artifacts.",
        "The project has datasets, metrics, model artifacts, and a next hardware path.",
        "Use this as an evidence board, not a bibliography page. Arrange sources and proof artifacts into clean groups.",
        [
            "Research base: VCTK, DNS Challenge, MAD military audio, ESC 50, AudioSet, RIRS NOISES, user recordings.",
            "Validation metrics: SNR, SI SNR, STOI, PESQ, speech deletion rate, active level error, latency.",
            "Artifacts: v12 metric model, v15 ONNX model, broad strict pass report, Pi 5 verification, failure analysis CSV.",
            "Next phase: synchronized dual microphone hardware and field recording fine tune.",
        ],
        [
            "Use a 2 by 2 evidence board: data, metrics, deployment, next hardware.",
            "Add small QR code placeholder if you will link demo/evidence later.",
            "Use one small final statement at bottom: AI engine validated, physical ANC hardware next.",
            "Keep reference names readable and do not paste long links on the slide.",
        ],
        "This slide should make judges feel that if they ask for proof, you have it ready. It should feel like a dossier index.",
        "Do not make a long citation dump. Use grouped evidence and keep detailed links in speaker notes or appendix if allowed.",
    )

    h(doc, "Creative demo concept for the PPT and video", 1)
    para(
        doc,
        "The demo should not look like a normal file player. It should look like a controlled field test. The phrase to use is controlled mission replay. "
        "The audio comes from local files, but the model processes it locally and frame by frame. This avoids unreliable laptop and Bluetooth hardware while staying honest.",
    )
    table(
        doc,
        ["Demo beat", "What appears on screen", "Why it feels believable"],
        [
            ["Mission select", "Operator chooses helicopter, vehicle, gunfire, siren, or fan scenario.", "Looks like a test harness, not a media player."],
            ["Incoming feed", "Noisy waveform and spectrogram scroll from left to right.", "Movement creates live feeling."],
            ["AI engine active", "ONNX backend, frame latency, model hash, and edge readiness update.", "Shows real local processing."],
            ["A B switch", "Noisy and enhanced audio toggle with same timeline.", "Judges hear the difference immediately."],
            ["Evidence snapshot", "SNR, STOI, PESQ, speech deletion and attenuation appear after scenario.", "Connects demo to measurable proof."],
        ],
        [1.4, 3.0, 2.5],
    )

    h(doc, "The six slide deck at a glance", 1)
    table(
        doc,
        ["Slide", "One second identity", "Five second proof"],
        [
            ["1", "Defence speech problem", "Team@ has a serious mission communication idea."],
            ["2", "Solution engine", "Noisy audio enters, enhanced speech and telemetry exit."],
            ["3", "Engineering pipeline", "Datasets, model, ONNX, dashboard and edge path exist."],
            ["4", "Feasibility", "Export, broad gate and Pi 5 benchmark are passed, with hardware risk understood."],
            ["5", "Impact", "Clearer commands, local processing, speech protection and field deployment path."],
            ["6", "Evidence", "Sources, metrics, artifacts and next hardware phase are ready."],
        ],
        [0.65, 2.25, 4.1],
    )

    h(doc, "What to remove from the final deck", 1)
    bullets(doc, [
        "Do not mention every model version. Keep v12 and v15 only unless asked.",
        "Do not show failed training logs or raw command outputs.",
        "Do not include long PS description text.",
        "Do not over-explain Bluetooth failure in the main flow. Mention it as hardware risk only.",
        "Do not claim final physical ANC headset. Claim AI engine and edge-ready prototype.",
        "Do not use more than one paragraph per slide.",
    ])

    h(doc, "Exact final positioning", 1)
    table(
        doc,
        ["Situation", "Say this"],
        [
            ["One line pitch", "VoxShield protects voice commands in defence noise and runs locally on edge hardware."],
            ["Demo label", "Controlled real time mission replay with local ONNX processing."],
            ["Current status", "AI enhancement engine validated, physical dual microphone ANC is the next hardware phase."],
            ["Why it beats ordinary AI projects", "Team@ built the model, the evaluation system, the dashboard, the edge benchmark and the failure analysis."],
            ["Limitation", "Military vehicle and helicopter STOI reliability remains the hardest case and guides the next data collection phase."],
        ],
        [1.6, 5.3],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
